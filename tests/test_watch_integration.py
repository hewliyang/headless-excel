import asyncio
import urllib.error
import urllib.request

import pytest
import pytest_asyncio
import websockets
from openpyxl import Workbook, load_workbook

from headless_excel.watch import watch


@pytest.fixture
def xlsx_file(tmp_path):
    path = tmp_path / "test.xlsx"
    wb = Workbook()
    wb.active["A1"] = "test"
    wb.save(path)
    wb.close()
    return path


@pytest_asyncio.fixture
async def watch_server(xlsx_file):
    http_port = 18080
    ws_port = 18765

    task = asyncio.create_task(
        watch(str(xlsx_file), http_port=http_port, ws_port=ws_port)
    )

    await asyncio.sleep(0.3)

    yield {"http_port": http_port, "ws_port": ws_port, "file": xlsx_file}

    task.cancel()
    try:
        await task
    except asyncio.CancelledError:
        pass


@pytest.mark.asyncio
async def test_http_server_serves_html(watch_server):
    port = watch_server["http_port"]

    loop = asyncio.get_event_loop()

    def fetch():
        with urllib.request.urlopen(f"http://localhost:{port}/", timeout=5) as resp:
            return resp.status, resp.read().decode()

    status, body = await loop.run_in_executor(None, fetch)

    assert status == 200
    assert "<html>" in body.lower()
    assert "test.xlsx" in body


@pytest.mark.asyncio
async def test_http_server_serves_xlsx(watch_server):
    port = watch_server["http_port"]

    loop = asyncio.get_event_loop()

    def fetch():
        with urllib.request.urlopen(
            f"http://localhost:{port}/file.xlsx", timeout=5
        ) as resp:
            return resp.status, resp.headers.get("Content-Type")

    status, content_type = await loop.run_in_executor(None, fetch)

    assert status == 200
    assert "spreadsheet" in content_type


@pytest.mark.asyncio
async def test_websocket_connects(watch_server):
    port = watch_server["ws_port"]

    async with websockets.connect(f"ws://localhost:{port}", close_timeout=1) as ws:
        assert ws is not None


@pytest.mark.asyncio
async def test_websocket_receives_reload_on_file_change(watch_server):
    port = watch_server["ws_port"]
    xlsx_file = watch_server["file"]

    async with websockets.connect(f"ws://localhost:{port}", close_timeout=1) as ws:
        await asyncio.sleep(0.1)

        wb = load_workbook(xlsx_file)
        wb.active["A1"] = "modified"
        wb.save(xlsx_file)
        wb.close()

        try:
            msg = await asyncio.wait_for(ws.recv(), timeout=2.0)
            assert msg == "reload"
        except asyncio.TimeoutError:
            pytest.fail("Did not receive reload message after file change")


@pytest.mark.asyncio
async def test_http_404_for_unknown_path(watch_server):
    port = watch_server["http_port"]

    loop = asyncio.get_event_loop()

    def fetch():
        try:
            urllib.request.urlopen(f"http://localhost:{port}/unknown", timeout=5)
            return 200
        except urllib.error.HTTPError as e:
            return e.code

    status = await loop.run_in_executor(None, fetch)
    assert status == 404


def test_watch_rejects_missing_file(tmp_path):
    missing = tmp_path / "missing.xlsx"

    with pytest.raises(FileNotFoundError):
        asyncio.run(watch(str(missing)))


def test_watch_rejects_non_xlsx(tmp_path):
    csv_file = tmp_path / "test.csv"
    csv_file.write_text("a,b,c")

    with pytest.raises(ValueError, match=".xlsx"):
        asyncio.run(watch(str(csv_file)))
