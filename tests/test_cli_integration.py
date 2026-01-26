import subprocess
import sys

from openpyxl import load_workbook


def run_cli(
    *args, input: str | None = None, timeout: int = 30
) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, "-m", "headless_excel.cli", *args],
        capture_output=True,
        text=True,
        input=input,
        timeout=timeout,
    )


class TestCheckCommand:
    def test_check_runs(self):
        result = run_cli("check")

        assert "headless-excel environment check" in result.stdout
        assert "LibreOffice" in result.stdout


class TestCreateCommand:
    def test_create_new_file(self, tmp_path):
        test_file = tmp_path / "new.xlsx"

        result = run_cli("create", str(test_file))

        assert result.returncode == 0
        assert test_file.exists()
        assert "Created" in result.stdout

    def test_create_file_is_valid_xlsx(self, tmp_path):
        test_file = tmp_path / "new.xlsx"
        run_cli("create", str(test_file))

        wb = load_workbook(test_file)
        assert len(wb.sheetnames) >= 1
        wb.close()


class TestEvalCommand:
    def test_eval_with_inline_code(self, tmp_path):
        test_file = tmp_path / "test.xlsx"
        run_cli("create", str(test_file))

        result = run_cli("eval", str(test_file), "ctx.active['A1'] = 42")

        assert result.returncode == 0

    def test_eval_from_stdin(self, tmp_path):
        test_file = tmp_path / "test.xlsx"
        run_cli("create", str(test_file))

        result = run_cli("eval", str(test_file), input="ctx.active['A1'] = 123")

        assert result.returncode == 0

    def test_eval_can_use_numberformats(self, tmp_path):
        test_file = tmp_path / "test.xlsx"
        run_cli("create", str(test_file))

        result = run_cli(
            "eval",
            str(test_file),
            "ctx.active['A1'].number_format = NumberFormats.ACCOUNTING",
        )

        assert result.returncode == 0

    def test_eval_file_not_found(self, tmp_path):
        missing_file = tmp_path / "missing.xlsx"

        result = run_cli("eval", str(missing_file), "pass")

        assert result.returncode != 0
