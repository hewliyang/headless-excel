from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    GradientFill,
    PatternFill,
    Protection,
    Side,
)

from headless_excel import (
    ErrorDetail,
    ExcelContext,
    FormulaError,
    SyncResult,
    create,
    run,
)
from headless_excel.proxy import CellProxy


class TestExcelContext:
    def test_create_new_workbook(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Hello"
            ctx.workbook.save(path)

        assert path.exists()

    def test_load_existing_workbook(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Hello"
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == "Hello"  # type: ignore[union-attr]

    def test_file_not_found(self, tmp_path: Path):
        path = tmp_path / "nonexistent.xlsx"
        with pytest.raises(FileNotFoundError):
            ExcelContext(path)

    def test_workbook_aliases(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            assert ctx.workbook is ctx.wb

    def test_sheet_operations(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.create_sheet("Data")
            ctx.sheet("Data")["A1"] = "Test"
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert "Data" in ctx.workbook.sheetnames
            assert ctx.sheet("Data")["A1"].value == "Test"  # type: ignore[union-attr]

    def test_values_before_sync_raises(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(RuntimeError, match="Call sync"):
                _ = ctx.values


class TestRun:
    def test_run_opens_existing(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ctx.active["A1"] = "Hello"
            ctx.workbook.save(path)

        with run(path, auto_sync=False) as ctx:
            assert ctx.active["A1"].value == "Hello"  # type: ignore[union-attr]

    def test_run_file_not_found(self, tmp_path: Path):
        path = tmp_path / "nonexistent.xlsx"
        with pytest.raises(FileNotFoundError):
            with run(path) as _:
                pass

    def test_run_auto_sync_false(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ctx.active["A1"] = 100
            ctx.workbook.save(path)


class TestCreate:
    def test_create_new_file(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ctx.active["A1"] = "Created"
            ctx.workbook.save(path)

        assert path.exists()

    def test_create_file_exists_raises(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ctx.workbook.save(path)

        with pytest.raises(FileExistsError, match="already exists"):
            with create(path) as ctx:
                pass

    def test_create_overwrite(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ctx.active["A1"] = "Original"
            ctx.workbook.save(path)

        with create(path, overwrite=True, auto_sync=False) as ctx:
            ctx.active["A1"] = "Replaced"
            ctx.workbook.save(path)

        with run(path) as ctx:
            assert ctx.active["A1"].value == "Replaced"  # type: ignore[union-attr]

    def test_create_manual_sync_prints_errors_on_exit(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=True) as ctx:
            ws = ctx.active
            ws["A1"] = 100
            ws["A2"] = "=A1+Sheet2!A1"

            result = ctx.sync(raise_on_errors=False)
            assert not result.success
            assert result.total_errors == 1

        captured = capsys.readouterr()
        assert "#NAME?" in captured.err
        assert path.exists()

    def test_create_verbose_errors_prints_to_stderr(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=True) as ctx:
            ws = ctx.active
            ws["A1"] = 100
            ws["A2"] = "=A1+Sheet2!A1"

        captured = capsys.readouterr()
        assert "SyncResult(success=False" in captured.err
        assert "#NAME?" in captured.err
        assert path.exists()

    def test_create_verbose_errors_false(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws = ctx.active
            ws["A1"] = 100
            ws["A2"] = "=A1+Sheet2!A1"

        captured = capsys.readouterr()
        assert "#NAME?" not in captured.err
        assert path.exists()


class TestSyncResult:
    def test_sync_result_success(self, tmp_path: Path):
        result = SyncResult(success=True, total_errors=0, errors={})
        assert result.success
        result.raise_on_errors()

    def test_sync_result_errors(self):
        result = SyncResult(
            success=False,
            total_errors=2,
            errors={"#REF!": ["Sheet1!A1", "Sheet1!B2"]},
        )
        assert not result.success
        assert result.total_errors == 2

        with pytest.raises(FormulaError) as exc_info:
            result.raise_on_errors()

        assert exc_info.value.total == 2
        assert "#REF!" in exc_info.value.errors


class TestFormulaError:
    def test_formula_error_str(self):
        err = FormulaError(
            errors={"#REF!": ["Sheet1!A1"], "#NAME?": ["Sheet1!B2"]},
            total=2,
        )
        s = str(err)
        assert "Formula errors (2):" in s
        assert "#REF!" in s
        assert "#NAME?" in s

    def test_formula_error_str_with_details(self):
        err = FormulaError(
            errors={"#DIV/0!": ["Sheet1!A3"]},
            total=1,
            error_details=[
                ErrorDetail(
                    location="Sheet1!A3",
                    error="#DIV/0!",
                    formula="=A1/A2",
                    neighbors={"Sheet1!A1": 10, "Sheet1!A2": 0},
                )
            ],
        )
        s = str(err)
        assert "Formula errors (1):" in s
        assert "Sheet1!A3: #DIV/0!" in s
        assert "formula==A1/A2" in s
        assert "inputs={" in s
        assert "Sheet1!A1=10" in s
        assert "Sheet1!A2=0" in s


class TestRangeProxy:
    def test_range_values_read(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 1
            ctx.active["B1"] = 2
            ctx.active["A2"] = 3
            ctx.active["B2"] = 4
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            r = ctx.active.range("A1:B2")
            assert r.values == [[1, 2], [3, 4]]

    def test_range_values_write(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:C2")
            r.values = [[1, 2, 3], [4, 5, 6]]
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 1  # type: ignore[union-attr]
            assert ctx.active["C2"].value == 6  # type: ignore[union-attr]

    def test_range_values_write_row_mismatch_lenient(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:B2")
            r.values = [[1, 2]]
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 1  # type: ignore[union-attr]
            assert ctx.active["B1"].value == 2  # type: ignore[union-attr]
            assert ctx.active["A2"].value is None  # type: ignore[union-attr]

        captured = capsys.readouterr()
        assert "[headless-excel] .values wrote to A1:B1" in captured.err
        assert "specified A1:B2" in captured.err
        assert "got 1×2" in captured.err

    def test_range_values_write_col_mismatch_lenient(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:C2")
            r.values = [[1, 2], [3, 4]]
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 1  # type: ignore[union-attr]
            assert ctx.active["B2"].value == 4  # type: ignore[union-attr]
            assert ctx.active["C1"].value is None  # type: ignore[union-attr]

        captured = capsys.readouterr()
        assert "[headless-excel] .values wrote to A1:B2" in captured.err
        assert "specified A1:C2" in captured.err
        assert "got 2×2" in captured.err

    def test_range_values_write_overflow(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:B2")
            r.values = [[1, 2, 3], [4, 5, 6], [7, 8, 9]]
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 1  # type: ignore[union-attr]
            assert ctx.active["C3"].value == 9  # type: ignore[union-attr]

        captured = capsys.readouterr()
        assert "[headless-excel] .values wrote to A1:C3" in captured.err
        assert "specified A1:B2" in captured.err
        assert "got 3×3" in captured.err

    def test_range_values_write_exact_match(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:B2")
            r.values = [[1, 2], [3, 4]]
            ctx.workbook.save(path)

        captured = capsys.readouterr()
        assert "[headless-excel] .values wrote to A1:B2 (2×2)" in captured.err
        assert "specified" not in captured.err

    def test_range_values_write_not_list(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:B2")
            with pytest.raises(ValueError, match="must be a 2D list"):
                r.values = "not a list"  # type: ignore[assignment]

    def test_range_values_write_row_not_list(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:B2")
            with pytest.raises(ValueError, match="Row 0 must be a list"):
                r.values = ["not", "nested"]  # type: ignore[assignment]

    def test_range_shape(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("B3:E6")
            assert r.shape == (4, 4)
            assert r.num_rows == 4
            assert r.num_cols == 4

    def test_range_single_cell(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("C5")
            assert r.shape == (1, 1)
            r.values = [[42]]
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["C5"].value == 42  # type: ignore[union-attr]

    def test_range_formulas(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 10
            ctx.active["B1"] = "=A1*2"
            ctx.active["A2"] = "=A1+5"
            ctx.active["B2"] = 30

            r = ctx.active.range("A1:B2")
            formulas = r.formulas
            assert formulas == {"B1": "=A1*2", "A2": "=A1+5"}

    def test_range_values_write_with_formulas(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active.range("A1:C2").values = [
                [10, 20, "=A1+B1"],
                [30, 40, "=A2+B2"],
            ]
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 10  # type: ignore[union-attr]
            assert ctx.active["B2"].value == 40  # type: ignore[union-attr]
            assert ctx.active.formulas == {"C1": "=A1+B1", "C2": "=A2+B2"}
            assert ctx.active.range("C1:C2").formulas == {
                "C1": "=A1+B1",
                "C2": "=A2+B2",
            }

    def test_range_invalid_reference(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="Invalid range reference"):
                ctx.active.range("not_a_range")

    def test_range_repr(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:B2")
            assert "RangeProxy" in repr(r)
            assert "A1:B2" in repr(r)

    def test_range_values_empty_data(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            r = ctx.active.range("A1:B2")
            r.values = []

        captured = capsys.readouterr()
        assert "[headless-excel] .values no data to write" in captured.err


class TestWorksheetWrite:
    def test_write_basic(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            result = ctx.active.write("A1", [[1, 2, 3], [4, 5, 6]])
            ctx.workbook.save(path)

            assert result == "A1:C2"

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 1  # type: ignore[union-attr]
            assert ctx.active["C2"].value == 6  # type: ignore[union-attr]

        captured = capsys.readouterr()
        assert "[headless-excel] .write() wrote to A1:C2 (2×3)" in captured.err

    def test_write_single_cell(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            result = ctx.active.write("B5", [[42]])
            ctx.workbook.save(path)

            assert result == "B5"

        with ExcelContext(path) as ctx:
            assert ctx.active["B5"].value == 42  # type: ignore[union-attr]

        captured = capsys.readouterr()
        assert "[headless-excel] .write() wrote to B5 (1×1)" in captured.err

    def test_write_with_formulas(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            result = ctx.active.write(
                "A1",
                [
                    [10, 20, "=A1+B1"],
                    [30, 40, "=A2+B2"],
                ],
            )
            ctx.workbook.save(path)

            assert result == "A1:C2"

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 10  # type: ignore[union-attr]
            assert ctx.active.formulas == {"C1": "=A1+B1", "C2": "=A2+B2"}

    def test_write_marks_dirty(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.sync()
            assert not ctx._dirty

            ctx.active.write("A1", [[1]])
            assert ctx._dirty

    def test_write_rejects_range(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="single cell anchor"):
                ctx.active.write("A1:B2", [[1, 2], [3, 4]])

    def test_write_rejects_non_list(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="must be a 2D list"):
                ctx.active.write("A1", "not a list")  # type: ignore[arg-type]

    def test_write_rejects_non_list_rows(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="Row 0 must be a list"):
                ctx.active.write("A1", ["not", "nested"])  # type: ignore[list-item]

    def test_write_empty_data(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            result = ctx.active.write("A1", [])

            assert result == "A1"

        captured = capsys.readouterr()
        assert "[headless-excel] .write() no data to write" in captured.err

    def test_write_jagged_arrays(self, tmp_path: Path, capsys):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            result = ctx.active.write(
                "A1",
                [
                    [1, 2, 3],
                    [4, 5],
                    [7, 8, 9, 10],
                ],
            )
            ctx.workbook.save(path)

            assert result == "A1:D3"

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value == 1  # type: ignore[union-attr]
            assert ctx.active["D3"].value == 10  # type: ignore[union-attr]
            assert ctx.active["C2"].value is None  # type: ignore[union-attr]

        captured = capsys.readouterr()
        assert "[headless-excel] .write() wrote to A1:D3 (3×4)" in captured.err


class TestRangeApplyStyle:
    def test_apply_font_to_single_cell(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Hello"
            ctx.active.range("A1").apply_style(font=Font(bold=True, color="FF0000"))
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            cell: CellProxy = ctx.active["A1"]  # type: ignore[assignment]
            assert cell.font.bold is True
            assert cell.font.color.rgb == "00FF0000"  # type: ignore[union-attr]

    def test_apply_font_to_range(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            for col in ["A", "B", "C"]:
                ctx.active[f"{col}1"] = f"Value {col}"
            ctx.active.range("A1:C1").apply_style(font=Font(italic=True))
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            for col in ["A", "B", "C"]:
                assert ctx.active[f"{col}1"].font.italic is True  # type: ignore[union-attr]

    def test_apply_fill(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Colored"
            fill = PatternFill(
                start_color="FFFF00", end_color="FFFF00", fill_type="solid"
            )
            ctx.active.range("A1").apply_style(fill=fill)
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].fill.start_color.rgb == "00FFFF00"  # type: ignore[union-attr]

    def test_apply_gradient_fill(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Gradient"
            gradient = GradientFill(stop=["FF0000", "0000FF"])
            ctx.active.range("A1").apply_style(gradient_fill=gradient)
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            cell: CellProxy = ctx.active["A1"]  # type: ignore[assignment]
            assert cell.fill.type == "linear"  # type: ignore[union-attr]
            assert len(cell.fill.stop) == 2  # type: ignore[union-attr,arg-type]
            assert cell.fill.stop[0].color.rgb == "00FF0000"  # type: ignore[union-attr,index]
            assert cell.fill.stop[1].color.rgb == "000000FF"  # type: ignore[union-attr,index]

    def test_apply_gradient_fill_to_range(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            for col in ["A", "B", "C"]:
                ctx.active[f"{col}1"] = f"Gradient {col}"
            gradient = GradientFill(stop=["00FF00", "FFFF00"], degree=90)
            ctx.active.range("A1:C1").apply_style(gradient_fill=gradient)
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            for col in ["A", "B", "C"]:
                cell: CellProxy = ctx.active[f"{col}1"]  # type: ignore[assignment]
                assert cell.fill.type == "linear"  # type: ignore[union-attr]
                assert cell.fill.degree == 90  # type: ignore[union-attr]
                assert cell.fill.stop[0].color.rgb == "0000FF00"  # type: ignore[union-attr,index]
                assert cell.fill.stop[1].color.rgb == "00FFFF00"  # type: ignore[union-attr,index]

    def test_gradient_fill_takes_precedence_over_fill(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Both fills"
            pattern = PatternFill(start_color="FF0000", fill_type="solid")
            gradient = GradientFill(stop=["00FF00", "0000FF"])
            ctx.active.range("A1").apply_style(fill=pattern, gradient_fill=gradient)
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            cell: CellProxy = ctx.active["A1"]  # type: ignore[assignment]
            assert cell.fill.type == "linear"  # type: ignore[union-attr]
            assert len(cell.fill.stop) == 2  # type: ignore[union-attr,arg-type]

    def test_apply_number_format(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 1234.56
            ctx.active["A2"] = -789.12
            ctx.active.range("A1:A2").apply_style(number_format="#,##0.00;(#,##0.00)")
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].number_format == "#,##0.00;(#,##0.00)"  # type: ignore[union-attr]
            assert ctx.active["A2"].number_format == "#,##0.00;(#,##0.00)"  # type: ignore[union-attr]

    def test_apply_alignment(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Centered"
            ctx.active.range("A1").apply_style(
                alignment=Alignment(horizontal="center", vertical="top")
            )
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].alignment.horizontal == "center"  # type: ignore[union-attr]
            assert ctx.active["A1"].alignment.vertical == "top"  # type: ignore[union-attr]

    def test_apply_border(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Bordered"
            thin_side = Side(style="thin", color="000000")
            border = Border(
                left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
            )
            ctx.active.range("A1").apply_style(border=border)
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].border.left.style == "thin"  # type: ignore[union-attr]
            assert ctx.active["A1"].border.right.style == "thin"  # type: ignore[union-attr]

    def test_apply_protection(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Protected"
            ctx.active.range("A1").apply_style(
                protection=Protection(locked=True, hidden=True)
            )
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].protection.locked is True  # type: ignore[union-attr]
            assert ctx.active["A1"].protection.hidden is True  # type: ignore[union-attr]

    def test_apply_multiple_styles(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["B2"] = 42
            ctx.active.range("B2").apply_style(
                font=Font(bold=True),
                fill=PatternFill(start_color="00FF00", fill_type="solid"),
                number_format="0.00",
                alignment=Alignment(horizontal="right"),
            )
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            cell: CellProxy = ctx.active["B2"]  # type: ignore[assignment]
            assert cell.font.bold is True
            assert cell.fill.start_color.rgb == "0000FF00"  # type: ignore[union-attr]
            assert cell.number_format == "0.00"
            assert cell.alignment.horizontal == "right"

    def test_apply_style_to_2d_range(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            for row in range(1, 4):
                for col in ["A", "B", "C"]:
                    ctx.active[f"{col}{row}"] = f"{col}{row}"
            ctx.active.range("A1:C3").apply_style(font=Font(size=14))
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            for row in range(1, 4):
                for col in ["A", "B", "C"]:
                    assert ctx.active[f"{col}{row}"].font.size == 14  # type: ignore[union-attr]


class TestRangeClear:
    def test_clear_values(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active.range("A1:C3").values = [
                [1, 2, 3],
                [4, 5, 6],
                [7, 8, 9],
            ]
            ctx.active.range("A1:C3").clear()

            values = ctx.active.range("A1:C3").values
            assert values == [
                [None, None, None],
                [None, None, None],
                [None, None, None],
            ]

    def test_clear_preserves_styles_when_styles_false(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Test"
            ctx.active.range("A1").apply_style(font=Font(bold=True))
            ctx.active.range("A1").clear(styles=False)
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert ctx.active["A1"].value is None  # type: ignore[union-attr]
            assert ctx.active["A1"].font.bold is True  # type: ignore[union-attr]

    def test_clear_with_styles(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Test"
            ctx.active.range("A1").apply_style(
                font=Font(bold=True),
                fill=PatternFill(start_color="FF0000", fill_type="solid"),
                number_format="0.00",
            )
            ctx.active.range("A1").clear(styles=True)
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            cell: CellProxy = ctx.active["A1"]  # type: ignore[assignment]
            assert cell.value is None
            assert cell.font.bold is not True
            assert cell.number_format == "General"

    def test_clear_formulas(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 10
            ctx.active["A2"] = "=A1*2"
            ctx.active.range("A2").clear()

            assert ctx.active["A2"].formula is None  # type: ignore[union-attr]
            assert ctx.active["A2"].value is None  # type: ignore[union-attr]

    def test_clear_marks_dirty(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 10
            ctx.sync()
            assert not ctx._dirty

            ctx.active.range("A1").clear()
            assert ctx._dirty

    def test_clear_unmerges_merged_cells(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "Merged Title"
            ws._formula_ws.merge_cells("A1:C1")
            ws["A2"] = "Data"

            ws.range("A1:C2").clear()

            assert ws.cell(1, 1).value is None
            assert ws.cell(2, 1).value is None

            assert len(list(ws._formula_ws.merged_cells.ranges)) == 0  # type: ignore[arg-type]

    def test_clear_unmerges_only_overlapping_ranges(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "Merge 1"
            ws._formula_ws.merge_cells("A1:B1")
            ws["D1"] = "Merge 2"
            ws._formula_ws.merge_cells("D1:E1")

            ws.range("A1:B2").clear()

            assert ws.cell(1, 1).value is None

            merged_ranges = [str(r) for r in ws._formula_ws.merged_cells.ranges]  # type: ignore[union-attr]
            assert "D1:E1" in merged_ranges
            assert "A1:B1" not in merged_ranges

    def test_clear_partial_overlap_unmerges(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "Wide Merge"
            ws._formula_ws.merge_cells("A1:D1")

            ws.range("B1:C1").clear()

            assert len(list(ws._formula_ws.merged_cells.ranges)) == 0  # type: ignore[arg-type]

    def test_clear_with_styles_unmerges(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "Styled Merge"
            ws._formula_ws.merge_cells("A1:C1")
            ws.range("A1").apply_style(font=Font(bold=True))

            ws.range("A1:C1").clear(styles=True)

            assert ws.cell(1, 1).value is None
            assert ws.cell(1, 1).font.bold is not True
            assert len(list(ws._formula_ws.merged_cells.ranges)) == 0  # type: ignore[arg-type]


class TestAutoFill:
    def test_auto_fill_formula_down(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=B1*C1"
            ctx.active.range("A1:A5").auto_fill()

            assert ctx.active.get_cell("A1").formula == "=B1*C1"
            assert ctx.active.get_cell("A2").formula == "=B2*C2"
            assert ctx.active.get_cell("A3").formula == "=B3*C3"
            assert ctx.active.get_cell("A4").formula == "=B4*C4"
            assert ctx.active.get_cell("A5").formula == "=B5*C5"

    def test_auto_fill_formula_right(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=A2+A3"
            ctx.active.range("A1:D1").auto_fill(direction="right")

            assert ctx.active.get_cell("A1").formula == "=A2+A3"
            assert ctx.active.get_cell("B1").formula == "=B2+B3"
            assert ctx.active.get_cell("C1").formula == "=C2+C3"
            assert ctx.active.get_cell("D1").formula == "=D2+D3"

    def test_auto_fill_formula_up(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A5"] = "=B5*2"
            ctx.active.range("A1:A5").auto_fill(direction="up")

            assert ctx.active.get_cell("A5").formula == "=B5*2"
            assert ctx.active.get_cell("A4").formula == "=B4*2"
            assert ctx.active.get_cell("A3").formula == "=B3*2"
            assert ctx.active.get_cell("A2").formula == "=B2*2"
            assert ctx.active.get_cell("A1").formula == "=B1*2"

    def test_auto_fill_formula_left(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["D1"] = "=D2*2"
            ctx.active.range("A1:D1").auto_fill(direction="left")

            assert ctx.active.get_cell("D1").formula == "=D2*2"
            assert ctx.active.get_cell("C1").formula == "=C2*2"
            assert ctx.active.get_cell("B1").formula == "=B2*2"
            assert ctx.active.get_cell("A1").formula == "=A2*2"

    def test_auto_fill_absolute_references(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=$B$1*C1"
            ctx.active.range("A1:A3").auto_fill()

            assert ctx.active.get_cell("A1").formula == "=$B$1*C1"
            assert ctx.active.get_cell("A2").formula == "=$B$1*C2"
            assert ctx.active.get_cell("A3").formula == "=$B$1*C3"

    def test_auto_fill_mixed_references(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=$B1*C$1"
            ctx.active.range("A1:A3").auto_fill()

            assert ctx.active.get_cell("A1").formula == "=$B1*C$1"
            assert ctx.active.get_cell("A2").formula == "=$B2*C$1"
            assert ctx.active.get_cell("A3").formula == "=$B3*C$1"

    def test_auto_fill_values(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 100
            ctx.active.range("A1:A5").auto_fill()

            for row in range(1, 6):
                assert ctx.active.get_cell(f"A{row}").value == 100

    def test_auto_fill_cross_sheet_reference(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=Sheet2!A1+B1"
            ctx.active.range("A1:A3").auto_fill()

            assert ctx.active.get_cell("A1").formula == "=Sheet2!A1+B1"
            assert ctx.active.get_cell("A2").formula == "=Sheet2!A2+B2"
            assert ctx.active.get_cell("A3").formula == "=Sheet2!A3+B3"

    def test_auto_fill_sum_range(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=SUM(B1:D1)"
            ctx.active.range("A1:A3").auto_fill()

            assert ctx.active.get_cell("A1").formula == "=SUM(B1:D1)"
            assert ctx.active.get_cell("A2").formula == "=SUM(B2:D2)"
            assert ctx.active.get_cell("A3").formula == "=SUM(B3:D3)"

    def test_auto_fill_multi_row_source(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=X1"
            ctx.active["A2"] = "=Y1"
            ctx.active.range("A1:A8").auto_fill(source_rows=2)

            assert ctx.active.get_cell("A1").formula == "=X1"
            assert ctx.active.get_cell("A2").formula == "=Y1"
            assert ctx.active.get_cell("A3").formula == "=X3"
            assert ctx.active.get_cell("A4").formula == "=Y3"
            assert ctx.active.get_cell("A5").formula == "=X5"
            assert ctx.active.get_cell("A6").formula == "=Y5"
            assert ctx.active.get_cell("A7").formula == "=X7"
            assert ctx.active.get_cell("A8").formula == "=Y7"

    def test_auto_fill_multi_col_source(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=A2"
            ctx.active["B1"] = "=B3"
            ctx.active.range("A1:F1").auto_fill(direction="right", source_cols=2)

            assert ctx.active.get_cell("A1").formula == "=A2"
            assert ctx.active.get_cell("B1").formula == "=B3"
            assert ctx.active.get_cell("C1").formula == "=C2"
            assert ctx.active.get_cell("D1").formula == "=D3"
            assert ctx.active.get_cell("E1").formula == "=E2"
            assert ctx.active.get_cell("F1").formula == "=F3"

    def test_auto_fill_copies_styles(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 100
            ctx.active.get_cell("A1").font = Font(bold=True, color="0000FF")
            ctx.active.get_cell("A1").number_format = "#,##0.00"
            ctx.active.range("A1:A3").auto_fill()

            for row in range(1, 4):
                cell = ctx.active._formula_ws.cell(row, 1)
                assert cell.font.bold is True
                assert cell.font.color.rgb == "000000FF"
                assert cell.number_format == "#,##0.00"

    def test_auto_fill_skip_styles(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 100
            ctx.active.get_cell("A1").font = Font(bold=True)
            ctx.active.range("A1:A3").auto_fill(copy_styles=False)

            assert ctx.active._formula_ws.cell(1, 1).font.bold is True
            assert ctx.active._formula_ws.cell(2, 1).font.bold is False
            assert ctx.active._formula_ws.cell(3, 1).font.bold is False

    def test_auto_fill_single_cell_raises(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="Cannot auto_fill a single cell"):
                ctx.active.range("A1").auto_fill()

    def test_auto_fill_source_rows_too_large_raises(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="source_rows.*must be less than"):
                ctx.active.range("A1:A3").auto_fill(source_rows=3)

    def test_auto_fill_source_cols_too_large_raises(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="source_cols.*must be less than"):
                ctx.active.range("A1:C1").auto_fill(direction="right", source_cols=3)

    def test_auto_fill_auto_direction_tall(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=B1"
            ctx.active.range("A1:A5").auto_fill()

            assert ctx.active.get_cell("A2").formula == "=B2"

    def test_auto_fill_auto_direction_wide(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=A2"
            ctx.active.range("A1:E1").auto_fill()

            assert ctx.active.get_cell("B1").formula == "=B2"

    def test_auto_fill_marks_dirty(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 100
            ctx.sync()
            assert not ctx._dirty

            ctx.active.range("A1:A3").auto_fill()
            assert ctx._dirty

    def test_auto_fill_after_sync_calculates_values(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["B1"] = 10
            ctx.active["C1"] = 2
            ctx.active["C2"] = 3
            ctx.active["C3"] = 4
            ctx.sync()

            ctx.active["D1"] = "=$B$1*C1"
            ctx.active.range("D1:D3").auto_fill()
            ctx.sync()

            assert ctx.active["D1"].value == 20  # type: ignore[union-attr]
            assert ctx.active["D2"].value == 30  # type: ignore[union-attr]
            assert ctx.active["D3"].value == 40  # type: ignore[union-attr]

    def test_formula_after_sync_calculates(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 10
            ctx.active["B1"] = 5
            ctx.sync()

            ctx.active["C1"] = "=A1+B1"
            ctx.sync()

            assert ctx.active["C1"].value == 15  # type: ignore[union-attr]


class TestSheetFormulas:
    def test_sheet_formulas(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 10
            ctx.active["A2"] = 20
            ctx.active["A3"] = "=A1+A2"
            ctx.active["B1"] = "=A1*2"

            formulas = ctx.active.formulas
            assert formulas == {"A3": "=A1+A2", "B1": "=A1*2"}

    def test_sheet_formulas_empty(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            formulas = ctx.active.formulas
            assert formulas == {}

    def test_sheet_formulas_no_formulas(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = 100
            ctx.active["A2"] = "text"
            formulas = ctx.active.formulas
            assert formulas == {}

    def test_sheet_formulas_complex(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=SUM(B1:B10)"
            ctx.active["A2"] = "=IF(A1>0,A1,-A1)"
            ctx.active["A3"] = "=VLOOKUP(A1,B:C,2,FALSE)"

            formulas = ctx.active.formulas
            assert "A1" in formulas
            assert "A2" in formulas
            assert "A3" in formulas


class TestExtractCellRefs:
    def test_simple_refs(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            refs = ctx._extract_cell_refs("=A1+B2", "Sheet1")
            assert set(refs) == {"Sheet1!A1", "Sheet1!B2"}

    def test_absolute_refs(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            refs = ctx._extract_cell_refs("=$A$1+$B2+C$3", "Sheet1")
            assert "Sheet1!A1" in refs
            assert "Sheet1!B2" in refs
            assert "Sheet1!C3" in refs

    def test_cross_sheet_refs(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            refs = ctx._extract_cell_refs("=Sheet2!A1+B1", "Sheet1")
            assert "Sheet2!A1" in refs
            assert "Sheet1!B1" in refs

    def test_range_refs(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            refs = ctx._extract_cell_refs("=SUM(A1:A10)", "Sheet1")
            assert "Sheet1!A1" in refs
            assert "Sheet1!A10" in refs

    def test_no_refs(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            refs = ctx._extract_cell_refs("=PI()", "Sheet1")
            assert refs == []

    def test_mixed_content(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            refs = ctx._extract_cell_refs('=IF(A1>0, "YES", B2)', "Sheet1")
            assert "Sheet1!A1" in refs
            assert "Sheet1!B2" in refs


class TestErrorDetail:
    def test_error_detail_creation(self):
        detail = ErrorDetail(
            location="Sheet1!A1",
            error="#DIV/0!",
            formula="=B1/C1",
            neighbors={"Sheet1!B1": 100, "Sheet1!C1": 0},
        )
        assert detail.location == "Sheet1!A1"
        assert detail.error == "#DIV/0!"
        assert detail.formula == "=B1/C1"
        assert detail.neighbors["Sheet1!B1"] == 100
        assert detail.neighbors["Sheet1!C1"] == 0

    def test_error_detail_defaults(self):
        detail = ErrorDetail(location="Sheet1!A1", error="#REF!")
        assert detail.formula is None
        assert detail.neighbors == {}


class TestSyncResultWithDetails:
    def test_sync_result_with_error_details(self):
        details = [
            ErrorDetail(
                location="Sheet1!A1",
                error="#DIV/0!",
                formula="=B1/C1",
                neighbors={"Sheet1!B1": 10, "Sheet1!C1": 0},
            )
        ]
        result = SyncResult(
            success=False,
            total_errors=1,
            errors={"#DIV/0!": ["Sheet1!A1"]},
            error_details=details,
        )
        assert len(result.error_details) == 1
        assert result.error_details[0].formula == "=B1/C1"
        assert result.error_details[0].neighbors["Sheet1!C1"] == 0

    def test_sync_result_empty_details(self):
        result = SyncResult(success=True, total_errors=0, errors={})
        assert result.error_details == []


class TestProxyValueUpdate:
    def test_proxy_shows_calculated_value_after_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ws = ctx.create_sheet("Sheet1", 0)

            ws["A1"] = 100
            ws["A2"] = 200
            ws["A3"] = "=SUM(A1:A2)"

            assert ws["A3"].value == "=SUM(A1:A2)"  # type: ignore[union-attr]

            ctx.sync()

            assert ws["A3"].value == 300  # type: ignore[union-attr]
            assert ws.range("A3").values == [[300]]

    def test_proxy_identity_maintained_after_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ws_before = ctx.active
            ws_before["A1"] = "=10+20"

            ctx.sync()

            ws_after = ctx.active

            assert ws_before is ws_after
            assert ws_after["A1"].value == 30  # type: ignore[union-attr]

    def test_create_sheet_proxy_updates_after_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ws = ctx.create_sheet("Data", 0)

            ws.range("A1:A3").values = [[100], [200], ["=A1+A2"]]

            assert ws["A3"].value == "=A1+A2"  # type: ignore[union-attr]

            ctx.sync()

            assert ws["A3"].value == 300  # type: ignore[union-attr]
            assert ws.range("A3").values == [[300]]

    def test_renamed_sheet_proxy_updates_after_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with create(path, auto_sync=False) as ctx:
            ws = ctx.active
            ws.title = "Renamed"

            ws["A1"] = 100
            ws["A2"] = "=A1*2"

            assert ws["A2"].value == "=A1*2"  # type: ignore[union-attr]

            ctx.sync()

            assert ws["A2"].value == 200  # type: ignore[union-attr]
            assert ws.range("A1:A2").values == [[100], [200]]


class TestBuildErrorDetails:
    def test_build_error_details_requires_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "=1/0"
            ctx.workbook.save(path)
            details = ctx._build_error_details({"#DIV/0!": ["Sheet!A1"]})
            assert details == []

    def test_build_error_details_basic(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = 10
        ws["B1"] = 0
        ws["C1"] = "=A1/B1"
        wb.save(path)

        with ExcelContext(path) as ctx:
            ctx._workbook = wb
            ctx._values_workbook = Workbook()
            ctx._values_workbook.active["A1"] = 10
            ctx._values_workbook.active["B1"] = 0
            ctx._values_workbook.active["C1"] = "#DIV/0!"

            details = ctx._build_error_details({"#DIV/0!": ["Sheet!C1"]})

            assert len(details) == 1
            assert details[0].location == "Sheet!C1"
            assert details[0].error == "#DIV/0!"
            assert details[0].formula == "=A1/B1"
            assert "Sheet!A1" in details[0].neighbors
            assert "Sheet!B1" in details[0].neighbors

    def test_auto_remove_default_sheet(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            assert len(ctx.workbook.sheetnames) == 1
            assert "Sheet" in ctx.workbook.sheetnames

            ctx.create_sheet("Data")

            assert len(ctx.workbook.sheetnames) == 1
            assert "Data" in ctx.workbook.sheetnames
            assert "Sheet" not in ctx.workbook.sheetnames

            ctx.create_sheet("Summary")
            assert len(ctx.workbook.sheetnames) == 2
            assert "Data" in ctx.workbook.sheetnames
            assert "Summary" in ctx.workbook.sheetnames

    def test_keep_default_sheet_if_has_data(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.active["A1"] = "Important data"

            ctx.create_sheet("Data")

            assert len(ctx.workbook.sheetnames) == 2
            assert "Sheet" in ctx.workbook.sheetnames
            assert "Data" in ctx.workbook.sheetnames
            assert ctx.sheet("Sheet")["A1"].value == "Important data"  # type: ignore[union-attr]

    def test_no_auto_remove_for_loaded_workbooks(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.workbook.save(path)

        with ExcelContext(path) as ctx:
            assert "Sheet" in ctx.workbook.sheetnames

            ctx.create_sheet("Data")

            assert len(ctx.workbook.sheetnames) == 2
            assert "Sheet" in ctx.workbook.sheetnames
            assert "Data" in ctx.workbook.sheetnames


class TestSheetManagement:
    def test_set_active_by_proxy(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.create_sheet("Data")
            ctx.create_sheet("Summary")

            ctx.active = ctx.sheet("Summary")
            assert ctx.active.title == "Summary"

            ctx.active = ctx.sheet("Data")
            assert ctx.active.title == "Data"

    def test_set_active_by_name(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.create_sheet("Data")
            ctx.create_sheet("Summary")

            ctx.active = "Summary"
            assert ctx.active.title == "Summary"

            ctx.active = "Data"
            assert ctx.active.title == "Data"

    def test_set_active_via_workbook_proxy(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.create_sheet("Data")
            ctx.create_sheet("Summary")

            ctx.wb.active = ctx.wb["Summary"]
            assert ctx.wb.active.title == "Summary"

            ctx.wb.active = "Data"
            assert ctx.wb.active.title == "Data"

    def test_delete_sheet(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.create_sheet("Data")
            ctx.create_sheet("Summary")

            assert "Data" in ctx.workbook.sheetnames
            assert "Summary" in ctx.workbook.sheetnames

            ctx.delete_sheet("Data")

            assert "Data" not in ctx.workbook.sheetnames
            assert "Summary" in ctx.workbook.sheetnames

    def test_delete_sheet_not_found(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.create_sheet("Data")

            with pytest.raises(KeyError, match="NotFound"):
                ctx.delete_sheet("NotFound")

    def test_delete_only_sheet_raises(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            with pytest.raises(ValueError, match="Cannot delete the only sheet"):
                ctx.delete_sheet("Sheet")

    def test_delete_sheet_clears_cache(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ctx.create_sheet("Data")
            ctx.create_sheet("Summary")

            _ = ctx.sheet("Data")
            assert "Data" in ctx._proxy._sheet_cache  # type: ignore[union-attr]

            ctx.delete_sheet("Data")

            assert "Data" not in ctx._proxy._sheet_cache  # type: ignore[union-attr]


class TestCellFormula:
    def test_formula_returns_formula_string(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "=SUM(B1:B10)"
            ws["A2"] = "=A1*2"

            assert ws["A1"].formula == "=SUM(B1:B10)"  # type: ignore[union-attr]
            assert ws["A2"].formula == "=A1*2"  # type: ignore[union-attr]

    def test_formula_returns_none_for_values(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = 42
            ws["A2"] = "hello"
            ws["A3"] = 3.14

            assert ws["A1"].formula is None  # type: ignore[union-attr]
            assert ws["A2"].formula is None  # type: ignore[union-attr]
            assert ws["A3"].formula is None  # type: ignore[union-attr]

    def test_formula_returns_none_for_empty(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            assert ws["A1"].formula is None  # type: ignore[union-attr]


class TestRangeDump:
    def test_dump_returns_formatted_string(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws.range("A1:C3").values = [
                [1, 2, 3],
                [4, 5, 6],
                [7, 8, 9],
            ]

            result = ws.range("A1:C3").dump()

            assert isinstance(result, str)
            assert "A" in result
            assert "B" in result
            assert "C" in result
            assert "1" in result
            assert "9" in result

    def test_dump_shows_formulas(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["B1"] = "=A1*2"

            result = ws.range("A1:B1").dump(show_formulas=True)

            assert "=A1*2" in result

    def test_dump_handles_none_values(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = 1

            result = ws.range("A1:B1").dump()

            assert isinstance(result, str)
            assert "1" in result

    def test_dump_truncates_long_values(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "This is a very long string that should be truncated"

            result = ws.range("A1:A1").dump()

            assert "…" in result


class TestFormulaCount:
    def test_range_formula_count(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = "=A1*2"
            ws["A3"] = "=A1+A2"
            ws["B1"] = "text"
            ws["B2"] = "=A1+10"

            assert ws.range("A1:A3").formula_count == 2

            assert ws.range("A1:B2").formula_count == 2

            assert ws.range("A1:B3").formula_count == 3

    def test_range_formula_count_empty(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = "text"

            assert ws.range("A1:A2").formula_count == 0

    def test_sheet_formula_count(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = "=A1*2"
            ws["B5"] = "=SUM(A1:A2)"
            ws["C10"] = "=A1+B5"

            assert ws.formula_count == 3

    def test_sheet_formula_count_empty(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = "text"

            assert ws.formula_count == 0


class TestFindErrors:
    def test_range_find_errors_after_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = 0
            ws["A3"] = "=A1/A2"
            ws["B1"] = "=NonExistent()"

            ctx.sync()

            errors = ws.range("A1:A3").find_errors()
            assert "#DIV/0!" in errors
            assert "A3" in errors["#DIV/0!"]

            errors = ws.range("B1:B1").find_errors()
            assert "#NAME?" in errors
            assert "B1" in errors["#NAME?"]

    def test_range_find_errors_no_errors(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = 2
            ws["A3"] = "=A1/A2"

            ctx.sync()

            errors = ws.range("A1:A3").find_errors()
            assert errors == {}

    def test_range_find_errors_before_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "=1/0"

            errors = ws.range("A1:A1").find_errors()
            assert errors == {}

    def test_sheet_find_errors_after_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = 0
            ws["A3"] = "=A1/A2"
            ws["C10"] = "=MissingSheet!A1"

            ctx.sync()

            errors = ws.find_errors()
            assert "#DIV/0!" in errors
            assert "A3" in errors["#DIV/0!"]
            assert len(errors) >= 1

    def test_sheet_find_errors_no_errors(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = "=A1*2"

            ctx.sync()

            errors = ws.find_errors()
            assert errors == {}

    def test_context_find_errors_single_sheet(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws = ctx.active
            ws["A1"] = 0
            ws["A2"] = "=1/A1"

            ctx.sync()

            errors = ctx.find_errors()
            assert "#DIV/0!" in errors
            assert any("A2" in loc for loc in errors["#DIV/0!"])

    def test_context_find_errors_multiple_sheets(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws1 = ctx.create_sheet("Sheet1")
            ws1["A1"] = 0
            ws1["A2"] = "=1/A1"

            ws2 = ctx.create_sheet("Sheet2")
            ws2["B1"] = "=UnknownFunc()"

            ctx.sync()

            errors = ctx.find_errors()

            assert "#DIV/0!" in errors
            assert "#NAME?" in errors

            assert any("Sheet1!" in loc for loc in errors["#DIV/0!"])
            assert any("Sheet2!" in loc for loc in errors["#NAME?"])

    def test_context_find_errors_no_errors(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path, verbose_errors=False) as ctx:
            ws = ctx.active
            ws["A1"] = 10
            ws["A2"] = "=A1*2"

            ctx.sync()

            errors = ctx.find_errors()
            assert errors == {}

    def test_context_find_errors_before_sync(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with ExcelContext(path, create=True) as ctx:
            ws = ctx.active
            ws["A1"] = "=1/0"

            errors = ctx.find_errors()
            assert errors == {}


class TestSourceCode:
    def test_code_none_for_library_usage(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"

        with create(path) as ctx:
            assert ctx._code is None

    def test_code_with_explicit_param(self, tmp_path: Path):
        path = tmp_path / "test.xlsx"
        code = "ctx.active['A1'] = 1\nctx.active['A2'] = 2\n"

        with create(path, _code=code) as ctx:
            assert ctx._code == code
