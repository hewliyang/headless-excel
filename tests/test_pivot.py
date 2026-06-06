"""Tests for the fluent pivot-table builder.

Structural tests run without LibreOffice; the recompute test is marked so it can
be skipped where LibreOffice is unavailable.
"""

import openpyxl
import pytest

from headless_excel import PivotTable, create
from headless_excel.pivot import _resolve_func, _split_source
from headless_excel.proxy import WorksheetProxy


def _seed(ctx):
    d = ctx.active
    d.title = "Data"
    d.range("A1:D7").values = [
        ["Region", "Product", "Units", "Sales"],
        ["North", "A", 10, 100],
        ["North", "B", 5, 80],
        ["South", "A", 7, 70],
        ["South", "B", 3, 60],
        ["North", "A", 4, 40],
        ["South", "B", 9, 90],
    ]
    return d


def test_resolve_func_aliases():
    assert _resolve_func("avg") == "average"
    assert _resolve_func("Sum") == "sum"
    with pytest.raises(ValueError):
        _resolve_func("median")


def test_split_source():
    assert _split_source("Data!A1:D7", "X") == ("Data", "A1:D7")
    assert _split_source("'My Sheet'!A1", "X") == ("My Sheet", "A1")
    assert _split_source("A1:D7", "Active") == ("Active", "A1:D7")


def test_builder_structure(tmp_path):
    with create(str(tmp_path / "p.xlsx"), overwrite=True) as ctx:
        _seed(ctx)
        p = ctx.create_sheet("Pivot")
        pt = (
            p.add_pivot("Data!A1:D7", anchor="A3", name="PT")
            .rows("Region")
            .columns("Product")
            .values("Sales", func="sum", num_format="$#,##0")
        )
        assert isinstance(pt, PivotTable)
        table = p._formula_ws._pivots[0]
        assert table.name == "PT"
        assert len(table.pivotFields) == 4
        assert table.pivotFields[0].axis == "axisRow"
        assert table.pivotFields[1].axis == "axisCol"
        assert table.pivotFields[3].dataField is True
        assert [f.x for f in table.rowFields] == [0]
        assert [f.x for f in table.colFields] == [1]
        assert table.dataFields[0].fld == 3
        assert table.dataFields[0].subtotal == "sum"
        assert table.dataFields[0].numFmtId is not None
        # cache is empty + refreshOnLoad so the app computes it
        assert table.cache.refreshOnLoad is True
        assert table.cache.records.count == 0
        assert [cf.name for cf in table.cache.cacheFields] == [
            "Region",
            "Product",
            "Units",
            "Sales",
        ]


def test_unknown_field_raises(tmp_path):
    with create(str(tmp_path / "p.xlsx"), overwrite=True) as ctx:
        _seed(ctx)
        p = ctx.create_sheet("Pivot")
        with pytest.raises(ValueError):
            p.add_pivot("Data!A1:D7", name="PT").rows("Nope")


def test_declarative_form(tmp_path):
    with create(str(tmp_path / "p.xlsx"), overwrite=True) as ctx:
        _seed(ctx)
        p = ctx.create_sheet("Pivot")
        p.add_pivot(
            "Data!A1:D7",
            anchor="A3",
            name="PT",
            rows=["Region"],
            columns=["Product"],
            values=[("Sales", "sum"), ("Units", "average", "Avg Units")],
            style="PivotStyleMedium9",
        )
        table = p._formula_ws._pivots[0]
        assert len(table.dataFields) == 2
        assert table.dataFields[1].subtotal == "average"
        assert table.dataFields[1].name == "Avg Units"
        assert table.pivotTableStyleInfo.name == "PivotStyleMedium9"


def _build_two(ctx):
    _seed(ctx)
    p = ctx.create_sheet("Pivot")
    p.add_pivot(
        "Data!A1:D7",
        anchor="A3",
        name="PT",
        rows=["Region"],
        columns=["Product"],
        values=[("Sales", "sum", "Total", "$#,##0")],
        filters=["Units"],
        style="PivotStyleMedium9",
    )
    p.add_pivot(
        "Data!A1:D7", anchor="H3", name="PT2", rows=["Product"], values=["Sales"]
    )
    return p


# -- Read --------------------------------------------------------------------


def test_read_pivots_and_get(tmp_path):
    with create(str(tmp_path / "p.xlsx"), overwrite=True) as ctx:
        p = _build_two(ctx)
        assert [pt.name for pt in p.pivots] == ["PT", "PT2"]
        pt = p.get_pivot("PT")
        assert pt.row_fields == ["Region"]
        assert pt.column_fields == ["Product"]
        assert pt.filter_fields == ["Units"]
        assert pt.value_fields == [
            {
                "field": "Sales",
                "func": "sum",
                "display_name": "Total",
                "num_format": "$#,##0",
            }
        ]
        assert pt.style_name == "PivotStyleMedium9"
        assert pt.source == "Data!A1:D7"
        assert pt.fields == ["Region", "Product", "Units", "Sales"]
        with pytest.raises(KeyError):
            p.get_pivot("Missing")


def test_read_after_reload(tmp_path):
    path = tmp_path / "p.xlsx"
    with create(str(path), overwrite=True, auto_sync=False) as ctx:
        _build_two(ctx)
        ctx._workbook.save(str(path))

    wb = openpyxl.load_workbook(path)
    p = WorksheetProxy(wb["Pivot"])
    assert {pt.name for pt in p.pivots} == {"PT", "PT2"}
    snap = p.get_pivot("PT").to_dict()
    assert snap["rows"] == ["Region"]
    assert snap["values"][0]["func"] == "sum"
    assert snap["values"][0]["num_format"] == "$#,##0"


# -- Update ------------------------------------------------------------------


def test_update_remove_and_clear(tmp_path):
    with create(str(tmp_path / "p.xlsx"), overwrite=True) as ctx:
        p = _build_two(ctx)
        pt = p.get_pivot("PT")
        pt.remove_columns("Product").remove_filters("Units")
        assert pt.column_fields == []
        assert pt.filter_fields == []
        # underlying table reflects the change
        table = p._formula_ws._pivots[0]
        assert table.colFields is None or all(f.x != 1 for f in table.colFields)
        pt.clear_values().values("Units", "average", "Avg")
        assert [v["field"] for v in pt.value_fields] == ["Units"]
        assert pt.value_fields[0]["func"] == "average"


def test_update_rename_move_grand_totals(tmp_path):
    with create(str(tmp_path / "p.xlsx"), overwrite=True) as ctx:
        p = _build_two(ctx)
        pt = p.get_pivot("PT")
        pt.rename("Renamed").move("A40").set_grand_totals(rows=False, columns=False)
        assert pt.name == "Renamed"
        table = next(t for t in p._formula_ws._pivots if t.name == "Renamed")
        assert table.location.ref == "A40"
        assert table.rowGrandTotals is False
        assert table.colGrandTotals is False
        # duplicate names are rejected
        with pytest.raises(ValueError):
            p.get_pivot("PT2").rename("Renamed")


def test_update_persists_after_reload(tmp_path):
    path = tmp_path / "p.xlsx"
    with create(str(path), overwrite=True, auto_sync=False) as ctx:
        _build_two(ctx)
        ctx._workbook.save(str(path))

    wb = openpyxl.load_workbook(path)
    p = WorksheetProxy(wb["Pivot"])
    p.get_pivot("PT").rename("Renamed").remove_columns("Product")
    wb.save(str(path))

    wb2 = openpyxl.load_workbook(path)
    p2 = WorksheetProxy(wb2["Pivot"])
    assert {pt.name for pt in p2.pivots} == {"Renamed", "PT2"}
    assert p2.get_pivot("Renamed").column_fields == []


# -- Delete ------------------------------------------------------------------


def test_delete_pivot(tmp_path):
    with create(str(tmp_path / "p.xlsx"), overwrite=True) as ctx:
        p = _build_two(ctx)
        assert p.remove_pivot("PT2") is True
        assert [pt.name for pt in p.pivots] == ["PT"]
        assert p.remove_pivot("Nope") is False
        # delete via the object too
        p.get_pivot("PT").delete()
        assert p.pivots == []


def test_delete_drops_cache_after_reload(tmp_path):
    path = tmp_path / "p.xlsx"
    with create(str(path), overwrite=True, auto_sync=False) as ctx:
        _build_two(ctx)
        ctx._workbook.save(str(path))

    wb = openpyxl.load_workbook(path)
    p = WorksheetProxy(wb["Pivot"])
    assert p.remove_pivot("PT") is True
    wb.save(str(path))

    wb2 = openpyxl.load_workbook(path)
    p2 = WorksheetProxy(wb2["Pivot"])
    assert [pt.name for pt in p2.pivots] == ["PT2"]


@pytest.mark.integration
def test_libreoffice_computes_pivot(tmp_path):
    path = tmp_path / "p.xlsx"
    with create(str(path), overwrite=True) as ctx:
        _seed(ctx)
        p = ctx.create_sheet("Pivot")
        (
            p.add_pivot("Data!A1:D7", anchor="A3", name="PT")
            .rows("Region")
            .columns("Product")
            .values("Sales", func="sum")
        )
        ctx.sync()

    wb = openpyxl.load_workbook(path, data_only=True)
    vals = {
        row[0]: row[1:]
        for row in wb["Pivot"].iter_rows(min_row=4, max_row=8, values_only=True)
        if row[0]
    }
    # North: A=100+40=140, B=80, total 220 ; South: A=70, B=60+90=150, total 220
    assert vals["North"][:3] == (140, 80, 220)
    assert vals["South"][:3] == (70, 150, 220)
    assert vals["Total Result"][:3] == (210, 230, 440)
