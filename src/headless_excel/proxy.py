"""Proxy wrappers for ergonomic value access."""

from __future__ import annotations

import re
import sys
from collections.abc import Callable
from copy import copy
from typing import Any, Literal

from openpyxl.cell.cell import ERROR_CODES, Cell
from openpyxl.formula.translate import Translator
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    GradientFill,
    PatternFill,
    Protection,
)
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from headless_excel.errors import ErrorScanResult

EXCEL_ERRORS = ERROR_CODES + ("#SPILL!", "#CALC!")

# Type alias for write callback
OnWriteCallback = Callable[[], None] | None

# Pattern for parsing range references like "A1:B2" or "A1"
RANGE_PATTERN = re.compile(r"^([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$", re.IGNORECASE)


def _parse_range(ref: str) -> tuple[int, int, int, int]:
    """Parse range reference into (min_row, min_col, max_row, max_col).

    Args:
        ref: Range like "A1:B2" or single cell like "A1"

    Returns:
        Tuple of (min_row, min_col, max_row, max_col), 1-indexed

    Raises:
        ValueError: If ref is not a valid range format
    """
    match = RANGE_PATTERN.match(ref.upper())
    if not match:
        raise ValueError(f"Invalid range reference: {ref!r}")

    start_col, start_row, end_col, end_row = match.groups()
    min_col = column_index_from_string(start_col)
    min_row = int(start_row)

    if end_col and end_row:
        max_col = column_index_from_string(end_col)
        max_row = int(end_row)
    else:
        max_col = min_col
        max_row = min_row

    # Normalize so min <= max
    if min_row > max_row:
        min_row, max_row = max_row, min_row
    if min_col > max_col:
        min_col, max_col = max_col, min_col

    return min_row, min_col, max_row, max_col


class RangeProxy:
    """Proxy for a cell range supporting bulk read/write operations.

    Example:
        r = sheet.range("A1:C3")
        r.values = [[1, 2, 3], [4, 5, 6], [7, 8, 9]]
        r.apply_style(font=Font(bold=True))
    """

    def __init__(
        self,
        ws_proxy: WorksheetProxy,
        range_ref: str,
        on_write: OnWriteCallback = None,
    ) -> None:
        """Initialize range proxy.

        Args:
            ws_proxy: Parent WorksheetProxy
            range_ref: Range reference like "A1:B2" or "A1"
            on_write: Callback to invoke when write operations occur
        """
        self._ws = ws_proxy
        self._range_ref = range_ref
        self._on_write = on_write
        self._min_row, self._min_col, self._max_row, self._max_col = _parse_range(
            range_ref
        )

    @property
    def num_rows(self) -> int:
        """Number of rows in range."""
        return self._max_row - self._min_row + 1

    @property
    def num_cols(self) -> int:
        """Number of columns in range."""
        return self._max_col - self._min_col + 1

    @property
    def shape(self) -> tuple[int, int]:
        """Shape as (rows, cols)."""
        return (self.num_rows, self.num_cols)

    @property
    def values(self) -> list[list[Any]]:
        """Get 2D array of values from the range."""
        result = []
        for row_idx in range(self._min_row, self._max_row + 1):
            row_values = []
            for col_idx in range(self._min_col, self._max_col + 1):
                cell = self._ws.cell(row_idx, col_idx)
                row_values.append(cell.value)
            result.append(row_values)
        return result

    @values.setter
    def values(self, data: list[list[Any]]) -> None:
        """Set 2D array of values to the range.

        The range acts as an anchor - data shape determines actual write region.
        This is lenient by design to avoid LLM off-by-one errors with range specs.

        Feedback is printed to stderr showing actual write region:
        - If data matches range: [headless-excel] .values wrote to A1:D3 (3×4)
        - If data mismatches: [headless-excel] .values wrote to A1:D3 (specified A1:D10, got 3×4)

        Args:
            data: 2D list of values to write (shape determines actual region)

        Raises:
            ValueError: If data is not a 2D list structure
        """
        # Validate structure (but not dimensions)
        if not isinstance(data, list):
            raise ValueError("Data must be a 2D list")

        if len(data) == 0:
            # Empty data - nothing to write
            print(
                "[headless-excel] .values no data to write (empty list)",
                file=sys.stderr,
            )
            return

        # Validate all rows are lists and find max col width
        max_cols = 0
        for i, row in enumerate(data):
            if not isinstance(row, list):
                raise ValueError(f"Row {i} must be a list")
            max_cols = max(max_cols, len(row))

        actual_rows = len(data)
        actual_cols = max_cols

        # Calculate actual write region
        actual_max_row = self._min_row + actual_rows - 1
        actual_max_col = self._min_col + actual_cols - 1

        # Build actual range string
        start_cell = f"{get_column_letter(self._min_col)}{self._min_row}"
        end_cell = f"{get_column_letter(actual_max_col)}{actual_max_row}"
        if start_cell == end_cell:
            actual_range = start_cell
        else:
            actual_range = f"{start_cell}:{end_cell}"

        # Check if there's a mismatch
        specified_matches = (
            actual_rows == self.num_rows and actual_cols == self.num_cols
        )

        # Mark dirty before writing
        if self._on_write:
            self._on_write()

        # Write values
        for row_offset, row_data in enumerate(data):
            for col_offset, value in enumerate(row_data):
                row_idx = self._min_row + row_offset
                col_idx = self._min_col + col_offset
                self._ws._formula_ws.cell(row_idx, col_idx).value = value

        # Print feedback to stderr
        if specified_matches:
            print(
                f"[headless-excel] .values wrote to {actual_range} ({actual_rows}×{actual_cols})",
                file=sys.stderr,
            )
        else:
            print(
                f"[headless-excel] .values wrote to {actual_range} (specified {self._range_ref}, got {actual_rows}×{actual_cols})",
                file=sys.stderr,
            )

    @property
    def formulas(self) -> dict[str, str]:
        """Get formulas in range as {coordinate: formula}.

        Only cells containing formulas are included.
        More token-efficient than a 2D array for sparse formulas.

        Example:
            >>> ws.range("A1:C3").formulas
            {'B2': '=A1*2', 'C3': '=SUM(A1:B2)'}
        """
        result: dict[str, str] = {}
        for row_idx in range(self._min_row, self._max_row + 1):
            for col_idx in range(self._min_col, self._max_col + 1):
                cell = self._ws._formula_ws.cell(row_idx, col_idx)
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    result[cell.coordinate] = val
        return result

    @property
    def formula_count(self) -> int:
        """Count of formulas in this range.

        Example:
            >>> ws.range("A1:C3").formula_count
            5
        """
        return len(self.formulas)

    def find_errors(self) -> ErrorScanResult:
        """Find formula errors in this range.

        Scans the values (materialized after sync) for Excel error values
        like #DIV/0!, #REF!, #NAME?, etc.

        Returns:
            ErrorScanResult with errors organized by type and total count.
            Has a nice __repr__ with truncation for agent-friendly output.

        Example:
            >>> result = ws.range("A1:C10").find_errors()
            >>> print(result)
            Formula errors (3):
              #DIV/0!: A3, B5
              #REF!: C10
        """
        if self._ws._values_ws is None:
            return ErrorScanResult()

        error_details: dict[str, list[str]] = {err: [] for err in EXCEL_ERRORS}

        for row_idx in range(self._min_row, self._max_row + 1):
            for col_idx in range(self._min_col, self._max_col + 1):
                cell = self._ws._values_ws.cell(row_idx, col_idx)
                if cell.value is not None and isinstance(cell.value, str):
                    for err in EXCEL_ERRORS:
                        if err in cell.value:
                            error_details[err].append(cell.coordinate)
                            break

        # Filter to only include error types that have occurrences
        errors_by_type = {k: v for k, v in error_details.items() if v}
        total = sum(len(locs) for locs in errors_by_type.values())
        return ErrorScanResult(errors_by_type=errors_by_type, total_errors=total)

    def dump(self, show_formulas: bool = False) -> str:
        """Print range contents as formatted table for debugging.

        Args:
            show_formulas: If True, show formulas instead of values

        Returns:
            Formatted string representation of the range

        Example:
            >>> print(ws.range("A1:C3").dump())
            |     A |     B |     C |
            |-------|-------|-------|
            |   100 |   200 |   300 |
            |    10 |    20 |    30 |
            |     1 |     2 |     3 |
        """
        # Collect data
        if show_formulas:
            data = []
            for row_idx in range(self._min_row, self._max_row + 1):
                row_data = []
                for col_idx in range(self._min_col, self._max_col + 1):
                    cell = self._ws._formula_ws.cell(row_idx, col_idx)
                    row_data.append(cell.value)
                data.append(row_data)
        else:
            data = self.values

        # Calculate column widths
        col_letters = [
            get_column_letter(c) for c in range(self._min_col, self._max_col + 1)
        ]
        col_widths = [len(letter) for letter in col_letters]

        for row in data:
            for i, val in enumerate(row):
                str_val = "" if val is None else str(val)
                col_widths[i] = max(col_widths[i], len(str_val))

        # Cap column width at 20 for readability
        col_widths = [min(w, 20) for w in col_widths]

        # Build output
        lines = []

        # Header row with column letters
        header = (
            "|"
            + "|".join(
                f" {col_letters[i]:>{col_widths[i]}} " for i in range(len(col_letters))
            )
            + "|"
        )
        lines.append(header)

        # Separator
        sep = (
            "|"
            + "|".join("-" * (col_widths[i] + 2) for i in range(len(col_widths)))
            + "|"
        )
        lines.append(sep)

        # Data rows
        for row in data:
            row_strs = []
            for i, val in enumerate(row):
                if val is None:
                    str_val = ""
                elif isinstance(val, float):
                    str_val = f"{val:.4g}"
                else:
                    str_val = str(val)
                # Truncate if too long
                if len(str_val) > col_widths[i]:
                    str_val = str_val[: col_widths[i] - 1] + "…"
                row_strs.append(f" {str_val:>{col_widths[i]}} ")
            lines.append("|" + "|".join(row_strs) + "|")

        result = "\n".join(lines)
        return result

    def clear(self, styles: bool = True) -> None:
        """Clear cell values and styles in the range.

        Args:
            styles: If True (default), also reset cell styles to defaults

        Example:
            >>> ws.range("A1:C3").clear()  # Clear values and styles
            >>> ws.range("A1:C3").clear(styles=False)  # Clear values only
        """
        if self._on_write:
            self._on_write()

        # Unmerge any merged cell ranges that overlap with this range
        merged_to_remove = []
        for merged_range in self._ws._formula_ws.merged_cells.ranges:  # type: ignore[union-attr]
            # Check if merged range overlaps with clear range
            if (
                merged_range.min_row <= self._max_row
                and merged_range.max_row >= self._min_row
                and merged_range.min_col <= self._max_col
                and merged_range.max_col >= self._min_col
            ):
                merged_to_remove.append(str(merged_range))

        for merged_addr in merged_to_remove:
            self._ws._formula_ws.unmerge_cells(merged_addr)

        for row_idx in range(self._min_row, self._max_row + 1):
            for col_idx in range(self._min_col, self._max_col + 1):
                cell = self._ws._formula_ws.cell(row_idx, col_idx)
                cell.value = None
                if styles:
                    cell.font = Font()
                    cell.fill = PatternFill()
                    cell.alignment = Alignment()
                    cell.border = Border()
                    cell.protection = Protection()
                    cell.number_format = "General"

    def apply_style(
        self,
        font: Font | None = None,
        fill: PatternFill | None = None,
        gradient_fill: GradientFill | None = None,
        alignment: Alignment | None = None,
        border: Border | None = None,
        protection: Protection | None = None,
        number_format: str | None = None,
    ) -> None:
        """Apply styles to all cells in the range.

        Args:
            font: Font style to apply
            fill: Fill/background style (PatternFill)
            gradient_fill: Gradient fill (takes precedence over fill)
            alignment: Alignment style
            border: Border style
            protection: Protection settings
            number_format: Number format string
        """
        # Mark dirty before applying styles
        if self._on_write:
            self._on_write()

        for row_idx in range(self._min_row, self._max_row + 1):
            for col_idx in range(self._min_col, self._max_col + 1):
                cell = self._ws._formula_ws.cell(row_idx, col_idx)
                if font is not None:
                    cell.font = font
                if gradient_fill is not None:
                    cell.fill = gradient_fill
                elif fill is not None:
                    cell.fill = fill
                if alignment is not None:
                    cell.alignment = alignment
                if border is not None:
                    cell.border = border
                if protection is not None:
                    cell.protection = protection
                if number_format is not None:
                    cell.number_format = number_format

    def auto_fill(
        self,
        direction: Literal["down", "right", "up", "left"] | None = None,
        source_rows: int = 1,
        source_cols: int = 1,
        copy_styles: bool = True,
    ) -> None:
        """Fill the range by repeating source cells, adjusting formula references.

        Mimics Excel's auto-fill (drag handle) behavior:
        - Formulas have their relative references adjusted
        - Absolute references ($A$1) stay fixed
        - Mixed references ($A1, A$1) adjust partially
        - Values (numbers, text) are copied as-is
        - Styles are copied by default

        The source cells are determined by direction:
        - down: First row(s) are source, fill downward
        - right: First column(s) are source, fill rightward
        - up: Last row(s) are source, fill upward
        - left: Last column(s) are source, fill leftward
        - None: Auto-detect based on range shape (tall=down, wide=right)

        Args:
            direction: Fill direction, or None to auto-detect
            source_rows: Number of source rows to repeat (for down/up fill)
            source_cols: Number of source columns to repeat (for left/right fill)
            copy_styles: If True (default), copy cell styles along with values

        Raises:
            ValueError: If range is a single cell (nothing to fill)
            ValueError: If source size >= range size in fill direction

        Example:
            >>> ws["A1"] = "=B1*C1"
            >>> ws.range("A1:A10").auto_fill()  # Fills A2:A10 with adjusted formulas

            >>> ws.range("A1:C1").values = [[1, 2, 3]]
            >>> ws.range("A1:C10").auto_fill()  # Fills rows 2-10

            >>> ws["A1"] = "=A2+A3"
            >>> ws.range("A1:D1").auto_fill(direction="right")  # Fills B1:D1

            # Multi-row source pattern (like Excel alternating rows)
            >>> ws["A1"] = "Revenue"
            >>> ws["A2"] = "Expenses"
            >>> ws.range("A1:A8").auto_fill(source_rows=2)  # Repeats pattern
        """
        # Validate range has something to fill
        if self.num_rows == 1 and self.num_cols == 1:
            raise ValueError("Cannot auto_fill a single cell range")

        # Auto-detect direction if not specified
        if direction is None:
            if self.num_rows >= self.num_cols:
                direction = "down"
            else:
                direction = "right"

        # Validate source size
        if direction in ("down", "up") and source_rows >= self.num_rows:
            raise ValueError(
                f"source_rows ({source_rows}) must be less than range rows ({self.num_rows})"
            )
        if direction in ("left", "right") and source_cols >= self.num_cols:
            raise ValueError(
                f"source_cols ({source_cols}) must be less than range cols ({self.num_cols})"
            )

        if self._on_write:
            self._on_write()

        if direction == "down":
            self._fill_down(copy_styles, source_rows)
        elif direction == "right":
            self._fill_right(copy_styles, source_cols)
        elif direction == "up":
            self._fill_up(copy_styles, source_rows)
        elif direction == "left":
            self._fill_left(copy_styles, source_cols)
        else:
            raise ValueError(f"Invalid direction: {direction!r}")

    def _copy_cell_style(self, source_cell, target_cell) -> None:
        """Copy all style attributes from source to target cell."""
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)
        target_cell.number_format = source_cell.number_format

    def _translate_formula(self, formula: str, from_coord: str, to_coord: str) -> str:
        """Translate a formula from one cell to another, adjusting references."""
        translator = Translator(formula, from_coord)
        return translator.translate_formula(to_coord)

    def _fill_down(self, copy_styles: bool, source_rows: int = 1) -> None:
        """Fill downward - first row(s) are source, repeating as pattern."""
        # Target rows start after source rows
        target_start = self._min_row + source_rows

        for target_row in range(target_start, self._max_row + 1):
            # Calculate which source row to use (cycling through source rows)
            source_offset = (target_row - target_start) % source_rows
            source_row = self._min_row + source_offset

            for col_idx in range(self._min_col, self._max_col + 1):
                source_cell = self._ws._formula_ws.cell(source_row, col_idx)
                target_cell = self._ws._formula_ws.cell(target_row, col_idx)
                col_letter = get_column_letter(col_idx)

                # Copy value, translating formula if needed
                value = source_cell.value
                if isinstance(value, str) and value.startswith("="):
                    from_coord = f"{col_letter}{source_row}"
                    to_coord = f"{col_letter}{target_row}"
                    value = self._translate_formula(value, from_coord, to_coord)
                target_cell.value = value

                if copy_styles:
                    self._copy_cell_style(source_cell, target_cell)

    def _fill_right(self, copy_styles: bool, source_cols: int = 1) -> None:
        """Fill rightward - first column(s) are source, repeating as pattern."""
        # Target cols start after source cols
        target_start = self._min_col + source_cols

        for target_col in range(target_start, self._max_col + 1):
            # Calculate which source column to use (cycling through source cols)
            source_offset = (target_col - target_start) % source_cols
            source_col = self._min_col + source_offset
            source_col_letter = get_column_letter(source_col)
            target_col_letter = get_column_letter(target_col)

            for row_idx in range(self._min_row, self._max_row + 1):
                source_cell = self._ws._formula_ws.cell(row_idx, source_col)
                target_cell = self._ws._formula_ws.cell(row_idx, target_col)

                # Copy value, translating formula if needed
                value = source_cell.value
                if isinstance(value, str) and value.startswith("="):
                    from_coord = f"{source_col_letter}{row_idx}"
                    to_coord = f"{target_col_letter}{row_idx}"
                    value = self._translate_formula(value, from_coord, to_coord)
                target_cell.value = value

                if copy_styles:
                    self._copy_cell_style(source_cell, target_cell)

    def _fill_up(self, copy_styles: bool, source_rows: int = 1) -> None:
        """Fill upward - last row(s) are source, repeating as pattern."""
        # Target rows end before source rows
        target_end = self._max_row - source_rows

        for target_row in range(target_end, self._min_row - 1, -1):
            # Calculate which source row to use (cycling through source rows)
            source_offset = (target_end - target_row) % source_rows
            source_row = self._max_row - source_offset

            for col_idx in range(self._min_col, self._max_col + 1):
                source_cell = self._ws._formula_ws.cell(source_row, col_idx)
                target_cell = self._ws._formula_ws.cell(target_row, col_idx)
                col_letter = get_column_letter(col_idx)

                # Copy value, translating formula if needed
                value = source_cell.value
                if isinstance(value, str) and value.startswith("="):
                    from_coord = f"{col_letter}{source_row}"
                    to_coord = f"{col_letter}{target_row}"
                    value = self._translate_formula(value, from_coord, to_coord)
                target_cell.value = value

                if copy_styles:
                    self._copy_cell_style(source_cell, target_cell)

    def _fill_left(self, copy_styles: bool, source_cols: int = 1) -> None:
        """Fill leftward - last column(s) are source, repeating as pattern."""
        # Target cols end before source cols
        target_end = self._max_col - source_cols

        for target_col in range(target_end, self._min_col - 1, -1):
            # Calculate which source column to use (cycling through source cols)
            source_offset = (target_end - target_col) % source_cols
            source_col = self._max_col - source_offset
            source_col_letter = get_column_letter(source_col)
            target_col_letter = get_column_letter(target_col)

            for row_idx in range(self._min_row, self._max_row + 1):
                source_cell = self._ws._formula_ws.cell(row_idx, source_col)
                target_cell = self._ws._formula_ws.cell(row_idx, target_col)

                # Copy value, translating formula if needed
                value = source_cell.value
                if isinstance(value, str) and value.startswith("="):
                    from_coord = f"{source_col_letter}{row_idx}"
                    to_coord = f"{target_col_letter}{row_idx}"
                    value = self._translate_formula(value, from_coord, to_coord)
                target_cell.value = value

                if copy_styles:
                    self._copy_cell_style(source_cell, target_cell)

    def __repr__(self) -> str:
        return f"<RangeProxy '{self._ws._formula_ws.title}'!{self._range_ref}>"


class CellProxy:
    """Proxy for openpyxl Cell that returns materialized values after sync.

    Delegates to the formula cell for writes and attributes,
    but returns values from the data-only cell when available.
    """

    def __init__(
        self,
        formula_cell: Cell,
        values_cell: Cell | None = None,
        on_write: OnWriteCallback = None,
    ) -> None:
        self._formula_cell = formula_cell
        self._values_cell = values_cell
        self._on_write = on_write

    def _update_values_cell(self, values_cell: Cell | None) -> None:
        """Update the values cell reference (called after sync)."""
        self._values_cell = values_cell

    @property
    def value(self) -> Any:
        """Get materialized value if available, otherwise formula value."""
        if self._values_cell is not None:
            return self._values_cell.value
        return self._formula_cell.value

    @value.setter
    def value(self, val: Any) -> None:
        """Set value on formula cell."""
        if self._on_write:
            self._on_write()
        self._formula_cell.value = val

    @property
    def formula(self) -> str | None:
        """Get formula string if cell contains a formula, otherwise None.

        Example:
            >>> ws['A1'] = '=SUM(B1:B10)'
            >>> ws['A1'].formula
            '=SUM(B1:B10)'
            >>> ws['A2'] = 42
            >>> ws['A2'].formula
            None
        """
        val = self._formula_cell.value
        if isinstance(val, str) and val.startswith("="):
            return val
        return None

    @property
    def font(self) -> Font:
        """Get font style."""
        return self._formula_cell.font  # type: ignore[return-value]

    @font.setter
    def font(self, value: Font) -> None:
        """Set font style."""
        if self._on_write:
            self._on_write()
        self._formula_cell.font = value

    @property
    def fill(self) -> PatternFill | GradientFill:
        """Get fill style."""
        return self._formula_cell.fill  # type: ignore[return-value]

    @fill.setter
    def fill(self, value: PatternFill | GradientFill) -> None:
        """Set fill style."""
        if self._on_write:
            self._on_write()
        self._formula_cell.fill = value

    @property
    def border(self) -> Border:
        """Get border style."""
        return self._formula_cell.border  # type: ignore[return-value]

    @border.setter
    def border(self, value: Border) -> None:
        """Set border style."""
        if self._on_write:
            self._on_write()
        self._formula_cell.border = value

    @property
    def number_format(self) -> str:
        """Get number format string."""
        return self._formula_cell.number_format  # type: ignore[return-value]

    @number_format.setter
    def number_format(self, value: str) -> None:
        """Set number format string."""
        if self._on_write:
            self._on_write()
        self._formula_cell.number_format = value

    @property
    def alignment(self) -> Alignment:
        """Get alignment style."""
        return self._formula_cell.alignment  # type: ignore[return-value]

    @alignment.setter
    def alignment(self, value: Alignment) -> None:
        """Set alignment style."""
        if self._on_write:
            self._on_write()
        self._formula_cell.alignment = value

    @property
    def protection(self) -> Protection:
        """Get protection settings."""
        return self._formula_cell.protection  # type: ignore[return-value]

    @protection.setter
    def protection(self, value: Protection) -> None:
        """Set protection settings."""
        if self._on_write:
            self._on_write()
        self._formula_cell.protection = value

    # Forward all other attributes to formula cell
    def __getattr__(self, name: str) -> Any:
        return getattr(self._formula_cell, name)

    def __setattr__(self, name: str, value: Any) -> None:
        if name in ("_formula_cell", "_values_cell", "_on_write"):
            object.__setattr__(self, name, value)
        else:
            # Mark dirty for any attribute write (font, number_format, etc.)
            on_write = object.__getattribute__(self, "_on_write")
            if on_write:
                on_write()
            setattr(self._formula_cell, name, value)

    def __repr__(self) -> str:
        return f"<CellProxy {self._formula_cell.coordinate}>"


class WorksheetProxy:
    """Proxy for openpyxl Worksheet that returns CellProxy objects.

    Wraps cell access to return CellProxy instances that show
    materialized values after sync.
    """

    def __init__(
        self,
        formula_ws: Worksheet,
        values_ws: Worksheet | None = None,
        on_write: OnWriteCallback = None,
    ) -> None:
        self._formula_ws = formula_ws
        self._values_ws = values_ws
        self._on_write = on_write
        self._cell_cache: dict[str, CellProxy] = {}  # Cache cell proxies

    def _update_formula_ws(self, formula_ws: Worksheet) -> None:
        """Update the formula worksheet and refresh all cached cell proxies."""
        self._formula_ws = formula_ws
        # Update all cached cell proxies with new formula cells
        for coord, cell_proxy in self._cell_cache.items():
            formula_cell = formula_ws[coord]
            cell_proxy._formula_cell = formula_cell

    def _update_values_ws(self, values_ws: Worksheet | None) -> None:
        """Update the values worksheet and refresh all cached cell proxies."""
        self._values_ws = values_ws
        # Update all cached cell proxies
        for coord, cell_proxy in self._cell_cache.items():
            values_cell = values_ws[coord] if values_ws else None
            cell_proxy._update_values_cell(values_cell)

    def __getitem__(self, key: str) -> CellProxy | tuple[tuple[CellProxy, ...], ...]:
        """Get cell or range, wrapped in CellProxy."""
        formula_item = self._formula_ws[key]

        # Single cell access like ws["A1"]
        if isinstance(formula_item, Cell):
            # Return cached proxy if it exists
            if key in self._cell_cache:
                return self._cell_cache[key]

            values_cell = None
            if self._values_ws is not None:
                values_cell = self._values_ws[key]
            proxy = CellProxy(formula_item, values_cell, self._on_write)
            self._cell_cache[key] = proxy
            return proxy

        # Range access like ws["A1:B2"] - returns tuple of tuples of cells
        if isinstance(formula_item, tuple):
            if self._values_ws is not None:
                values_item = self._values_ws[key]
                return tuple(
                    tuple(
                        CellProxy(f_cell, v_cell, self._on_write)
                        for f_cell, v_cell in zip(f_row, v_row)
                    )
                    for f_row, v_row in zip(formula_item, values_item)
                )
            else:
                return tuple(
                    tuple(CellProxy(cell, None, self._on_write) for cell in row)
                    for row in formula_item
                )

        return formula_item

    def __setitem__(self, key: str, value: Any) -> None:
        """Set cell value."""
        if self._on_write:
            self._on_write()
        self._formula_ws[key] = value

    def get_cell(self, ref: str) -> CellProxy:
        """Get a single cell by reference string like 'A1'.

        This is a type-safe alternative to __getitem__ for single cell access.

        Args:
            ref: Cell reference like 'A1'

        Returns:
            CellProxy for the cell

        Raises:
            ValueError: If ref is a range (contains ':')
        """
        if ":" in ref:
            raise ValueError(
                f"get_cell() only accepts single cell refs, not ranges: {ref}"
            )

        result = self[ref]
        # At runtime, single cell access always returns CellProxy
        return result  # type: ignore[return-value]

    def cell(self, row: int, column: int, value: Any = None) -> CellProxy:
        """Access cell by row/column index."""
        if value is not None and self._on_write:
            self._on_write()
        formula_cell = self._formula_ws.cell(row, column, value)
        values_cell = None
        if self._values_ws is not None:
            values_cell = self._values_ws.cell(row, column)
        return CellProxy(formula_cell, values_cell, self._on_write)

    def iter_rows(
        self, min_row=None, max_row=None, min_col=None, max_col=None, values_only=False
    ):
        """Iterate over rows, yielding CellProxy objects."""
        formula_rows = self._formula_ws.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
            values_only=values_only,
        )

        if values_only:
            # If values_only, openpyxl returns raw values, not cells
            yield from formula_rows
            return

        if self._values_ws is not None:
            values_rows = self._values_ws.iter_rows(
                min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
            )
            for f_row, v_row in zip(formula_rows, values_rows):
                yield tuple(
                    CellProxy(f_cell, v_cell, self._on_write)
                    for f_cell, v_cell in zip(f_row, v_row)
                )
        else:
            for f_row in formula_rows:
                yield tuple(CellProxy(cell, None, self._on_write) for cell in f_row)

    def iter_cols(
        self, min_row=None, max_row=None, min_col=None, max_col=None, values_only=False
    ):
        """Iterate over columns, yielding CellProxy objects."""
        formula_cols = self._formula_ws.iter_cols(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
            values_only=values_only,
        )

        if values_only:
            yield from formula_cols
            return

        if self._values_ws is not None:
            values_cols = self._values_ws.iter_cols(
                min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
            )
            for f_col, v_col in zip(formula_cols, values_cols):
                yield tuple(
                    CellProxy(f_cell, v_cell, self._on_write)
                    for f_cell, v_cell in zip(f_col, v_col)
                )
        else:
            for f_col in formula_cols:
                yield tuple(CellProxy(cell, None, self._on_write) for cell in f_col)

    def range(self, ref: str) -> RangeProxy:
        """Get a range for bulk read/write operations.

        Args:
            ref: Range reference like "A1:C3" or single cell "A1"

        Returns:
            RangeProxy for bulk operations

        Example:
            r = sheet.range("A1:C3")
            r.values = [[1, 2, 3], [4, 5, 6], [7, 8, 9]]
            r.apply_style(font=Font(bold=True))
        """
        return RangeProxy(self, ref, self._on_write)

    def write(self, cell: str, data: list[list[Any]]) -> str:
        """Write 2D data starting at anchor cell. Returns actual range written.

        This is the simplest API for bulk writes - no range math required.
        The data shape determines the write region automatically.

        Feedback is printed to stderr: [headless-excel] .write() wrote to A1:D3 (3×4)

        Args:
            cell: Anchor cell like "A1" (top-left of write region)
            data: 2D list of values to write

        Returns:
            The actual range written (e.g., "A1:D3")

        Raises:
            ValueError: If cell is invalid or data is not a 2D list

        Example:
            >>> ws.write("A1", [
            ...     ['Company', 'Revenue', 'Profit'],
            ...     ['Acme', 1000, 200],
            ...     ['Beta', 2000, 400],
            ... ])
            [headless-excel] .write() wrote to A1:C3 (3×3)
            'A1:C3'
        """
        # Validate cell reference
        if ":" in cell:
            raise ValueError(f"write() takes a single cell anchor, not a range: {cell}")

        # Validate structure
        if not isinstance(data, list):
            raise ValueError("Data must be a 2D list")

        if len(data) == 0:
            print(
                "[headless-excel] .write() no data to write (empty list)",
                file=sys.stderr,
            )
            return cell

        # Validate all rows are lists and find max col width
        max_cols = 0
        for i, row in enumerate(data):
            if not isinstance(row, list):
                raise ValueError(f"Row {i} must be a list")
            max_cols = max(max_cols, len(row))

        if max_cols == 0:
            print(
                "[headless-excel] .write() no data to write (empty rows)",
                file=sys.stderr,
            )
            return cell

        # Parse anchor cell
        min_row, min_col, _, _ = _parse_range(cell)

        actual_rows = len(data)
        actual_cols = max_cols

        # Calculate actual write region
        actual_max_row = min_row + actual_rows - 1
        actual_max_col = min_col + actual_cols - 1

        # Build actual range string
        start_cell = f"{get_column_letter(min_col)}{min_row}"
        end_cell = f"{get_column_letter(actual_max_col)}{actual_max_row}"
        if start_cell == end_cell:
            actual_range = start_cell
        else:
            actual_range = f"{start_cell}:{end_cell}"

        # Mark dirty before writing
        if self._on_write:
            self._on_write()

        # Write values
        for row_offset, row_data in enumerate(data):
            for col_offset, value in enumerate(row_data):
                row_idx = min_row + row_offset
                col_idx = min_col + col_offset
                self._formula_ws.cell(row_idx, col_idx).value = value

        # Print feedback to stderr
        print(
            f"[headless-excel] .write() wrote to {actual_range} ({actual_rows}×{actual_cols})",
            file=sys.stderr,
        )

        return actual_range

    @property
    def formulas(self) -> dict[str, str]:
        """Get all formulas in this sheet as {coordinate: formula}.

        Example:
            >>> sheet.formulas
            {'A1': '=B1+C1', 'D5': '=SUM(A1:A4)'}
        """
        result: dict[str, str] = {}
        for row in self._formula_ws.iter_rows():
            for cell in row:
                if (
                    cell.value is not None
                    and isinstance(cell.value, str)
                    and cell.value.startswith("=")
                ):
                    result[cell.coordinate] = cell.value
        return result

    @property
    def formula_count(self) -> int:
        """Count of formulas in this sheet.

        Example:
            >>> sheet.formula_count
            42
        """
        return len(self.formulas)

    def find_errors(self) -> ErrorScanResult:
        """Find formula errors in this sheet.

        Scans the values worksheet (materialized values after sync) for
        Excel error values like #DIV/0!, #REF!, #NAME?, etc.

        Returns:
            ErrorScanResult with errors organized by type and total count.
            Has a nice __repr__ with truncation for agent-friendly output.

        Example:
            >>> result = sheet.find_errors()
            >>> print(result)
            Formula errors (3):
              #DIV/0!: A3, B5
              #REF!: C10
        """
        if self._values_ws is None:
            return ErrorScanResult()

        error_details: dict[str, list[str]] = {err: [] for err in EXCEL_ERRORS}

        for row in self._values_ws.iter_rows():
            for cell in row:
                if cell.value is not None and isinstance(cell.value, str):
                    for err in EXCEL_ERRORS:
                        if err in cell.value:
                            error_details[err].append(cell.coordinate)
                            break

        # Filter to only include error types that have occurrences
        errors_by_type = {k: v for k, v in error_details.items() if v}
        total = sum(len(locs) for locs in errors_by_type.values())
        return ErrorScanResult(errors_by_type=errors_by_type, total_errors=total)

    # Forward all other attributes to formula worksheet
    def __getattr__(self, name: str) -> Any:
        return getattr(self._formula_ws, name)

    def __setattr__(self, name: str, value: Any) -> None:
        if name in ("_formula_ws", "_values_ws", "_on_write", "_cell_cache"):
            object.__setattr__(self, name, value)
        else:
            # Mark dirty for worksheet attribute writes (like title)
            on_write = object.__getattribute__(self, "_on_write")
            if on_write:
                on_write()
            setattr(self._formula_ws, name, value)

    def __repr__(self) -> str:
        return f"<WorksheetProxy '{self._formula_ws.title}'>"


class WorkbookProxy:
    """Proxy for openpyxl Workbook that returns WorksheetProxy objects.

    Wraps worksheet access to return WorksheetProxy instances.
    """

    def __init__(
        self,
        formula_wb: Workbook,
        values_wb: Workbook | None = None,
        on_write: OnWriteCallback = None,
    ) -> None:
        self._formula_wb = formula_wb
        self._values_wb = values_wb
        self._on_write = on_write
        self._sheet_cache: dict[str, WorksheetProxy] = {}  # Cache sheet proxies

    def _update_formula_wb(self, formula_wb: Workbook) -> None:
        """Update the formula workbook and refresh all cached sheet proxies."""
        self._formula_wb = formula_wb
        # Update all cached sheet proxies with new formula worksheet
        # Use actual sheet title (not cache key) since title may have changed
        for sheet_proxy in self._sheet_cache.values():
            actual_title = sheet_proxy._formula_ws.title
            if actual_title in formula_wb.sheetnames:
                formula_ws = formula_wb[actual_title]
                sheet_proxy._update_formula_ws(formula_ws)

    def _update_values_wb(self, values_wb: Workbook | None) -> None:
        """Update the values workbook and refresh all cached sheet proxies."""
        self._values_wb = values_wb
        # Update all cached sheet proxies with new values worksheet
        # Use actual sheet title (not cache key) since title may have changed
        for sheet_proxy in self._sheet_cache.values():
            actual_title = sheet_proxy._formula_ws.title
            values_ws = None
            if values_wb is not None and actual_title in values_wb.sheetnames:
                values_ws = values_wb[actual_title]
            sheet_proxy._update_values_ws(values_ws)

    def __getitem__(self, key: str) -> WorksheetProxy:
        """Get worksheet by name."""
        # Return cached proxy if it exists
        if key in self._sheet_cache:
            return self._sheet_cache[key]

        formula_ws = self._formula_wb[key]
        values_ws = None
        if self._values_wb is not None and key in self._values_wb.sheetnames:
            values_ws = self._values_wb[key]
        proxy = WorksheetProxy(formula_ws, values_ws, self._on_write)
        self._sheet_cache[key] = proxy
        return proxy

    @property
    def active(self) -> WorksheetProxy:
        """Get active worksheet."""
        formula_ws = self._formula_wb.active
        if formula_ws is None:
            raise RuntimeError("No active worksheet")

        # Use sheet name for caching to maintain identity
        sheet_name = formula_ws.title
        if sheet_name in self._sheet_cache:
            return self._sheet_cache[sheet_name]

        values_ws = None
        if self._values_wb is not None and sheet_name in self._values_wb.sheetnames:
            values_ws = self._values_wb[sheet_name]
        proxy = WorksheetProxy(formula_ws, values_ws, self._on_write)
        self._sheet_cache[sheet_name] = proxy
        return proxy

    @active.setter
    def active(self, sheet: WorksheetProxy | str) -> None:
        """Set active worksheet by proxy or name."""
        if self._on_write:
            self._on_write()
        if isinstance(sheet, str):
            self._formula_wb.active = self._formula_wb[sheet]
            if self._values_wb is not None and sheet in self._values_wb.sheetnames:
                self._values_wb.active = self._values_wb[sheet]
        elif isinstance(sheet, WorksheetProxy):
            self._formula_wb.active = sheet._formula_ws
            if self._values_wb is not None and sheet._values_ws is not None:
                self._values_wb.active = sheet._values_ws
        else:
            raise TypeError(f"Expected WorksheetProxy or str, got {type(sheet)}")

    def create_sheet(
        self, title: str | None = None, index: int | None = None
    ) -> WorksheetProxy:
        """Create a new worksheet."""
        if self._on_write:
            self._on_write()
        formula_ws = self._formula_wb.create_sheet(title, index)
        # New sheets don't have values yet, but cache so sync() updates them
        proxy = WorksheetProxy(formula_ws, None, self._on_write)
        if formula_ws.title:
            self._sheet_cache[formula_ws.title] = proxy
        return proxy

    @property
    def worksheets(self) -> list[WorksheetProxy]:
        """Get list of all worksheets."""
        result = []
        for f_ws in self._formula_wb.worksheets:
            values_ws = None
            if self._values_wb is not None and f_ws.title in self._values_wb.sheetnames:
                values_ws = self._values_wb[f_ws.title]
            result.append(WorksheetProxy(f_ws, values_ws, self._on_write))
        return result

    # Forward all other attributes to formula workbook
    def __getattr__(self, name: str) -> Any:
        return getattr(self._formula_wb, name)

    def __setattr__(self, name: str, value: Any) -> None:
        if name in ("_formula_wb", "_values_wb", "_sheet_cache", "_on_write"):
            object.__setattr__(self, name, value)
        elif name == "active":
            # Use our custom setter for active
            type(self).active.fset(self, value)
        else:
            # Mark dirty for any workbook attribute write
            on_write = object.__getattribute__(self, "_on_write")
            if on_write:
                on_write()
            setattr(self._formula_wb, name, value)

    def __repr__(self) -> str:
        return f"<WorkbookProxy with {len(self._formula_wb.worksheets)} sheets>"
