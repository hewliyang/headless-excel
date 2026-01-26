"""Excel context manager with automatic sync and recalculation."""

from __future__ import annotations

import logging
import re
import sys
from collections.abc import Generator
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, TypeVar

from openpyxl import Workbook, load_workbook

from headless_excel.daemon import is_daemon_running, start_daemon
from headless_excel.errors import (
    ErrorDetail,
    ErrorScanResult,
    FormulaError,
    SyncError,
    get_max_errors_displayed,
)
from headless_excel.hooks import (
    run_on_exit_hooks,
    run_post_sync_hooks,
    run_pre_sync_hooks,
)
from headless_excel.libre import recalc
from headless_excel.proxy import WorkbookProxy, WorksheetProxy

logger = logging.getLogger(__name__)

T = TypeVar("T")

CELL_REF_PATTERN = re.compile(
    r"(?<![A-Za-z_])(?:([A-Za-z_][A-Za-z0-9_]*!)?)?\$?([A-Z]{1,3})\$?([1-9][0-9]*)(?![A-Za-z0-9_])"
)


@dataclass
class SyncResult:
    """Result of a sync operation.

    Attributes:
        success: Whether sync completed without formula errors
        total_errors: Number of formula errors found
        errors: Dict mapping error type to list of cell locations
        error_details: List of detailed error information (when available)
    """

    success: bool
    total_errors: int = 0
    errors: dict[str, list[str]] = field(default_factory=dict)
    error_details: list[ErrorDetail] = field(default_factory=list)

    def raise_on_errors(self) -> None:
        """Raise FormulaError if any errors were found."""
        if not self.success:
            raise FormulaError(
                errors=self.errors,
                total=self.total_errors,
                error_details=self.error_details,
            )

    def __repr__(self) -> str:
        """Format sync result with truncation to prevent context pollution."""
        if self.success:
            return "SyncResult(success=True)"

        max_display = get_max_errors_displayed()
        lines = [f"SyncResult(success=False, total_errors={self.total_errors}):"]
        displayed = 0
        truncated = 0

        if self.error_details:
            for detail in self.error_details:
                if displayed >= max_display:
                    truncated += 1
                    continue
                parts = [f"  {detail.location}: {detail.error}"]
                if detail.formula:
                    parts.append(f" formula={detail.formula}")
                if detail.neighbors:
                    neighbors_str = ", ".join(
                        f"{k}={v}" for k, v in detail.neighbors.items()
                    )
                    parts.append(f" inputs={{{neighbors_str}}}")
                lines.append("".join(parts))
                displayed += 1
        else:
            for error_type, locations in self.errors.items():
                for loc in locations:
                    if displayed >= max_display:
                        truncated += 1
                        continue
                    lines.append(f"  {error_type}: {loc}")
                    displayed += 1

        if truncated > 0:
            lines.append(f"  ... and {truncated} more error(s) (truncated)")

        return "\n".join(lines)

    __str__ = __repr__


class ExcelContext:
    """Context for Excel operations with automatic sync support.

    Similar to OfficeJS Excel.run() pattern - make changes, then sync()
    to save, recalculate, and reload materialized values.

    Example:
        with ExcelContext("model.xlsx") as ctx:
            ctx.active["A1"] = "=SUM(B1:B10)"

            result = ctx.sync()  # saves, recalcs via LibreOffice, reloads
            if not result.success:
                print(f"Errors: {result.errors}")

            # After sync, .value returns materialized results automatically
            print(ctx.active["A1"].value)  # 450 (computed value)
    """

    def __init__(
        self,
        path: str | Path,
        create: bool = False,
        recalc_timeout: int = 30,
        _code: str | None = None,
    ) -> None:
        """Initialize Excel context.

        Args:
            path: Path to the Excel file
            create: If True, create new workbook; if False, load existing
            recalc_timeout: Timeout in seconds for LibreOffice recalculation
            _code: The code being executed (for hooks, set by CLI eval)
        """
        self.path = Path(path)
        self._recalc_timeout = recalc_timeout
        self._workbook: Workbook | None = None
        self._values_workbook: Workbook | None = None
        self._proxy: WorkbookProxy | None = None
        self._dirty = False
        self._is_new_workbook = create
        self._first_sheet_created = False
        self._last_sync_result: SyncResult | None = None
        self.__code = _code

        if create:
            self._workbook = Workbook()
        elif self.path.exists():
            self._workbook = load_workbook(self.path)
            self._values_workbook = load_workbook(self.path, data_only=True)
        else:
            raise FileNotFoundError(f"File not found: {path}")

        self._proxy = WorkbookProxy(
            self._workbook, self._values_workbook, on_write=self._mark_dirty
        )

    def _mark_dirty(self) -> None:
        """Mark the context as dirty (called by proxies on write operations)."""
        self._dirty = True

    @property
    def workbook(self) -> WorkbookProxy:
        """The workbook proxy (returns materialized values after sync)."""
        if self._proxy is None:
            raise RuntimeError("Context not initialized")
        return self._proxy

    @property
    def wb(self) -> WorkbookProxy:
        """Alias for workbook."""
        return self.workbook

    @property
    def values(self) -> Workbook:
        """The raw workbook with materialized values (read after sync).

        Note: Prefer using workbook/active/sheet() which return proxied
        cells that automatically show materialized values after sync.
        This property is for advanced use cases.
        """
        if self._values_workbook is None:
            raise RuntimeError("Call sync() first to materialize values")
        return self._values_workbook

    @property
    def _code(self) -> str | None:
        """The Python code being executed (if available).

        Only set when using CLI eval command. Useful for hooks that
        want to inspect or lint the user's code.

        Example:
            @on_exit
            def warn_large_blocks(ctx):
                if ctx._code:
                    lines = [l for l in ctx._code.splitlines() if l.strip()]
                    if len(lines) > 20:
                        print("Warning: consider splitting")
        """
        return self.__code

    @property
    def active(self) -> WorksheetProxy:
        """The active worksheet (returns materialized values after sync)."""
        ws = self.workbook.active
        if ws is None:
            raise RuntimeError("No active worksheet")
        return ws

    @active.setter
    def active(self, sheet: WorksheetProxy | str) -> None:
        """Set active worksheet by proxy or name.

        Example:
            ctx.active = ctx.sheet("Revenue")
            ctx.active = "Revenue"  # Also works
        """
        self._dirty = True
        self.workbook.active = sheet

    def sheet(self, name: str) -> WorksheetProxy:
        """Get worksheet by name."""
        return self.workbook[name]

    def delete_sheet(self, name: str) -> None:
        """Delete a worksheet by name.

        Args:
            name: Name of the sheet to delete

        Raises:
            KeyError: If sheet doesn't exist
            ValueError: If trying to delete the only sheet
        """
        if self._workbook is None:
            raise RuntimeError("Context not initialized")

        if name not in self._workbook.sheetnames:
            raise KeyError(f"Sheet '{name}' not found")

        if len(self._workbook.sheetnames) == 1:
            raise ValueError("Cannot delete the only sheet in workbook")

        self._dirty = True
        del self._workbook[name]

        # Also remove from values workbook if it exists
        if (
            self._values_workbook is not None
            and name in self._values_workbook.sheetnames
        ):
            del self._values_workbook[name]

        # Clear from proxy cache if it exists
        if self._proxy is not None and name in self._proxy._sheet_cache:
            del self._proxy._sheet_cache[name]

    def create_sheet(self, title: str, index: int | None = None) -> WorksheetProxy:
        """Create a new worksheet.

        Note: When creating a new workbook, the default empty "Sheet" is
        automatically removed when you create your first custom sheet.
        """
        self._dirty = True

        # Auto-remove the default "Sheet" on first create_sheet() call
        # for new workbooks (created with create=True)
        if (
            self._is_new_workbook
            and not self._first_sheet_created
            and self._workbook is not None
        ):
            self._first_sheet_created = True

            # Check if we have the default single empty sheet
            if len(self._workbook.worksheets) == 1:
                default_sheet = self._workbook.worksheets[0]

                # Check if it's the default "Sheet" and is empty
                if default_sheet.title == "Sheet":
                    # Check if sheet is empty (no data in any cells)
                    is_empty = True
                    for row in default_sheet.iter_rows():
                        for cell in row:
                            if cell.value is not None:
                                is_empty = False
                                break
                        if not is_empty:
                            break

                    # Remove the default sheet if empty
                    if is_empty:
                        self._workbook.remove(default_sheet)

        return self.workbook.create_sheet(title, index)

    def sync(self, raise_on_errors: bool = False) -> SyncResult:
        """Save, recalculate via LibreOffice, and reload with materialized values.

        This is the core operation - similar to OfficeJS ctx.sync().

        Pre-sync and post-sync hooks are called automatically. Configure hooks
        by placing Python files in ./.headless-excel/hooks/ or ~/.headless-excel/hooks/.

        Args:
            raise_on_errors: If True, raise FormulaError when errors found

        Returns:
            SyncResult with success status and any formula errors

        Raises:
            SyncError: If save or recalc fails
            FormulaError: If raise_on_errors=True and formula errors found
        """
        if self._workbook is None:
            raise RuntimeError("Context not initialized")

        run_pre_sync_hooks(self)

        # Save current workbook
        try:
            self._workbook.save(self.path)
        except Exception as e:
            raise SyncError(f"Failed to save workbook: {e}") from e

        # Recalculate via LibreOffice (raises RecalcError on failure)
        recalc(self.path, timeout=self._recalc_timeout)

        # Reload formula workbook (to pick up any LibreOffice changes)
        self._workbook = load_workbook(self.path)

        # Load values workbook
        self._values_workbook = load_workbook(self.path, data_only=True)

        # Update proxy with reloaded workbooks (preserves object identity)
        if self._proxy:
            self._proxy._update_formula_wb(self._workbook)
            self._proxy._update_values_wb(self._values_workbook)
        else:
            self._proxy = WorkbookProxy(
                self._workbook, self._values_workbook, on_write=self._mark_dirty
            )

        self._dirty = False

        # Scrape for formula errors
        error_scan = self.find_errors()
        error_details = self._build_error_details(error_scan.errors_by_type)

        sync_result = SyncResult(
            success=error_scan.total_errors == 0,
            total_errors=error_scan.total_errors,
            errors=error_scan.errors_by_type,
            error_details=error_details,
        )

        self._last_sync_result = sync_result

        run_post_sync_hooks(self, sync_result)

        if raise_on_errors:
            sync_result.raise_on_errors()

        logger.info("Sync successful, changes saved")
        return sync_result

    def find_errors(self) -> ErrorScanResult:
        """Find formula errors across all sheets.

        Scans all worksheets for Excel error values like #DIV/0!, #REF!, etc.

        Returns:
            ErrorScanResult with errors organized by type and total count.
            Has a nice __repr__ with truncation for agent-friendly output.

        Example:
            >>> result = ctx.find_errors()
            >>> print(result)
            Formula errors (2):
              #DIV/0!: Sheet1!A3
              #REF!: Sheet1!C10
        """
        if self._proxy is None:
            return ErrorScanResult()

        all_errors: dict[str, list[str]] = {}

        for ws_proxy in self._proxy.worksheets:
            sheet_errors = ws_proxy.find_errors()
            for err_type, locations in sheet_errors.errors_by_type.items():
                qualified = [f"{ws_proxy.title}!{loc}" for loc in locations]
                if err_type in all_errors:
                    all_errors[err_type].extend(qualified)
                else:
                    all_errors[err_type] = qualified

        total = sum(len(locs) for locs in all_errors.values())
        return ErrorScanResult(errors_by_type=all_errors, total_errors=total)

    def _extract_cell_refs(self, formula: str, default_sheet: str) -> list[str]:
        """Extract cell references from a formula.

        Args:
            formula: Formula string (e.g., '=B1/C1' or '=Sheet2!A1+B1')
            default_sheet: Default sheet name for unqualified references

        Returns:
            List of fully qualified cell references (e.g., ['Sheet1!B1', 'Sheet1!C1'])
        """
        refs = []
        for match in CELL_REF_PATTERN.finditer(formula):
            sheet_prefix, col, row = match.groups()
            if sheet_prefix:
                sheet_name = sheet_prefix.rstrip("!")
            else:
                sheet_name = default_sheet
            refs.append(f"{sheet_name}!{col}{row}")
        return refs

    def _build_error_details(
        self, error_locations: dict[str, list[str]]
    ) -> list[ErrorDetail]:
        """Build detailed error information for each error location.

        Args:
            error_locations: Dict mapping error type to list of locations

        Returns:
            List of ErrorDetail with formula and neighbor information
        """
        if self._workbook is None or self._values_workbook is None:
            return []

        details = []
        for error_type, locations in error_locations.items():
            for location in locations:
                if "!" in location:
                    sheet_name, cell_ref = location.split("!", 1)
                else:
                    sheet_name = (
                        self._workbook.active.title
                        if self._workbook.active
                        else "Sheet1"
                    )
                    cell_ref = location

                formula = None
                neighbors: dict[str, Any] = {}

                try:
                    formula_ws = self._workbook[sheet_name]
                    formula_cell = formula_ws[cell_ref]
                    if (
                        formula_cell.value
                        and isinstance(formula_cell.value, str)
                        and formula_cell.value.startswith("=")
                    ):
                        formula = formula_cell.value

                        cell_refs = self._extract_cell_refs(formula, sheet_name)
                        for ref in cell_refs:
                            ref_sheet, ref_cell = ref.split("!", 1)
                            try:
                                val_ws = self._values_workbook[ref_sheet]
                                neighbors[ref] = val_ws[ref_cell].value
                            except (KeyError, AttributeError):
                                neighbors[ref] = None
                except (KeyError, AttributeError):
                    pass

                details.append(
                    ErrorDetail(
                        location=location,
                        error=error_type,
                        formula=formula,
                        neighbors=neighbors,
                    )
                )

        return details

    def close(self) -> None:
        """Close both workbooks."""
        if self._workbook:
            self._workbook.close()
            self._workbook = None
        if self._values_workbook:
            self._values_workbook.close()
            self._values_workbook = None

    def __enter__(self) -> ExcelContext:
        return self

    def __exit__(self, exc_type: Any, exc_val: Any, exc_tb: Any) -> None:
        self.close()


def _try_start_daemon() -> None:
    """Try to start the LibreOffice daemon if not running.

    Silently ignores failures - recalc() will fall back to cold start.
    """
    try:
        if not is_daemon_running():
            start_daemon(wait=True, timeout=15)
    except Exception:
        pass


def _run_context(
    path: str | Path,
    create_mode: bool,
    auto_sync: bool,
    verbose_errors: bool,
    recalc_timeout: int,
    _code: str | None = None,
) -> Generator[ExcelContext, None, None]:
    """Internal context manager for Excel operations.

    Hooks are called automatically:
    - pre_sync: before each sync()
    - post_sync: after each sync()
    - on_exit: when context manager exits

    Automatically starts the LibreOffice daemon if not running.
    The daemon auto-exits after 5 minutes of inactivity.
    """
    _try_start_daemon()

    ctx = ExcelContext(
        path,
        create=create_mode,
        recalc_timeout=recalc_timeout,
        _code=_code,
    )
    try:
        yield ctx

        # Always sync on exit - detecting all mutations is impractical
        # Todo: rm ctx._dirty altogether
        if auto_sync:
            ctx.sync(raise_on_errors=False)

        # Print errors from any sync (manual or auto) to stderr
        # File is saved successfully - these are just formula errors in cells
        if (
            verbose_errors
            and ctx._last_sync_result
            and not ctx._last_sync_result.success
        ):
            for line in str(ctx._last_sync_result).split("\n"):
                print(f"[post-sync:errors] {line}", file=sys.stderr)

        run_on_exit_hooks(ctx)
    finally:
        ctx.close()


@contextmanager
def run(
    path: str | Path,
    auto_sync: bool = True,
    verbose_errors: bool = True,
    recalc_timeout: int = 30,
    _code: str | None = None,
) -> Generator[ExcelContext, None, None]:
    """Open an existing Excel file for operations.

    Similar to OfficeJS Excel.run() - provides a context, handles cleanup,
    and optionally auto-syncs on exit.

    Hooks are called automatically at key points:
    - pre_sync: before each sync() (e.g., auto-formatting)
    - post_sync: after each sync() (e.g., logging)
    - on_exit: when context exits (e.g., linting)

    Configure hooks by placing Python files in:
    - ./.headless-excel/hooks/ (project-local, takes precedence)
    - ~/.headless-excel/hooks/ (global fallback)

    Example:
        with run("model.xlsx") as ctx:
            ctx.active["A1"] = 100
            ctx.active["A2"] = "=A1*2"
            # auto-syncs on exit, hooks run automatically

        # Or with explicit sync:
        with run("model.xlsx", auto_sync=False) as ctx:
            ctx.active["A1"] = 100
            ctx.sync()  # manual sync
            print(ctx.values.active["A1"].value)

    Args:
        path: Path to existing Excel file
        auto_sync: If True, automatically sync on context exit
        verbose_errors: If True, print formula errors to stderr on exit.
            File is always saved - these are just errors in cell values.
        recalc_timeout: Timeout in seconds for LibreOffice recalculation
        _code: Internal use - source code string for CLI eval

    Yields:
        ExcelContext for operations

    Raises:
        FileNotFoundError: If file does not exist
    """
    yield from _run_context(
        path,
        False,
        auto_sync,
        verbose_errors,
        recalc_timeout,
        _code,
    )


@contextmanager
def create(
    path: str | Path,
    overwrite: bool = False,
    auto_sync: bool = True,
    verbose_errors: bool = True,
    recalc_timeout: int = 30,
    _code: str | None = None,
) -> Generator[ExcelContext, None, None]:
    """Create a new Excel file.

    Hooks are called automatically at key points:
    - pre_sync: before each sync() (e.g., auto-formatting)
    - post_sync: after each sync() (e.g., logging)
    - on_exit: when context exits (e.g., linting)

    Configure hooks by placing Python files in:
    - ./.headless-excel/hooks/ (project-local, takes precedence)
    - ~/.headless-excel/hooks/ (global fallback)

    Example:
        with create("new.xlsx") as ctx:
            ctx.active["A1"] = "Hello"
            # auto-syncs on exit, hooks run automatically

        # Overwrite existing file:
        with create("existing.xlsx", overwrite=True) as ctx:
            ctx.active["A1"] = "Fresh start"

    Args:
        path: Path for the new Excel file
        overwrite: If True, overwrite existing file; if False, raise FileExistsError
        auto_sync: If True, automatically sync on context exit
        verbose_errors: If True, print formula errors to stderr on exit.
            File is always saved - these are just errors in cell values.
        recalc_timeout: Timeout in seconds for LibreOffice recalculation
        _code: Internal use - source code string for CLI eval

    Yields:
        ExcelContext for operations

    Raises:
        FileExistsError: If file exists and overwrite=False
    """
    p = Path(path)
    if p.exists() and not overwrite:
        raise FileExistsError(
            f"File already exists: {path}. Use overwrite=True to replace."
        )
    yield from _run_context(
        path,
        True,
        auto_sync,
        verbose_errors,
        recalc_timeout,
        _code,
    )
