#!/usr/bin/env python3
"""
Benchmark current openpyxl + LibreOffice approach.

Run: uv run python scripts/openpyxl_benchmark.py
"""

import tempfile
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook


def benchmark_openpyxl():
    """Benchmark openpyxl operations."""
    print("=== openpyxl Benchmark ===\n")

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = Path(f.name)

    try:
        # Test 1: Create and write
        print("1. Create workbook + write 100 cells...")
        t0 = time.perf_counter()
        wb = Workbook()
        ws = wb.active
        for i in range(100):
            ws.cell(row=i + 1, column=1, value=i + 1)
        ws["B1"] = "=SUM(A1:A100)"
        create_time = time.perf_counter() - t0
        print(f"   Create + write: {create_time*1000:.2f}ms")

        # Test 2: Save
        print("\n2. Save to disk...")
        t0 = time.perf_counter()
        wb.save(path)
        save_time = time.perf_counter() - t0
        print(f"   Save: {save_time*1000:.2f}ms")
        wb.close()

        # Test 3: Load (formulas only)
        print("\n3. Load workbook (formulas)...")
        t0 = time.perf_counter()
        wb = load_workbook(path)
        load_formula_time = time.perf_counter() - t0
        print(f"   Load (formulas): {load_formula_time*1000:.2f}ms")
        print(f"   B1 value: {wb.active['B1'].value}")  # Shows formula, not result
        wb.close()

        # Test 4: Load (data_only) - gets cached values
        print("\n4. Load workbook (data_only)...")
        t0 = time.perf_counter()
        wb = load_workbook(path, data_only=True)
        load_values_time = time.perf_counter() - t0
        print(f"   Load (data_only): {load_values_time*1000:.2f}ms")
        print(f"   B1 value: {wb.active['B1'].value}")  # None - no cached value yet!
        wb.close()

        # Test 5: Double load (what we currently do)
        print("\n5. Double load (current approach)...")
        t0 = time.perf_counter()
        wb_formula = load_workbook(path)
        wb_values = load_workbook(path, data_only=True)
        double_load_time = time.perf_counter() - t0
        print(f"   Double load: {double_load_time*1000:.2f}ms")
        wb_formula.close()
        wb_values.close()

        # Test 6: Full sync cycle (without LibreOffice recalc)
        print("\n6. Full write-save-double_load cycle...")
        wb = load_workbook(path)
        t0 = time.perf_counter()
        wb.active["A1"] = 1000  # Modify
        wb.save(path)
        wb.close()
        wb_formula = load_workbook(path)
        wb_values = load_workbook(path, data_only=True)
        full_cycle_time = time.perf_counter() - t0
        print(f"   Full cycle (no recalc): {full_cycle_time*1000:.2f}ms")
        wb_formula.close()
        wb_values.close()

        # Test 7: With LibreOffice recalc
        print("\n7. Full sync with LibreOffice recalc...")
        from headless_excel.daemon import is_daemon_running
        from headless_excel.libre import recalc

        wb = load_workbook(path)
        ws = wb.active
        ws["A1"] = 1000

        t0 = time.perf_counter()
        wb.save(path)
        save_t = time.perf_counter() - t0

        t1 = time.perf_counter()
        recalc(path)
        recalc_t = time.perf_counter() - t1

        t2 = time.perf_counter()
        wb_formula = load_workbook(path)
        wb_values = load_workbook(path, data_only=True)
        reload_t = time.perf_counter() - t2

        total = time.perf_counter() - t0

        print(f"   Save:        {save_t*1000:.2f}ms")
        print(
            f"   Recalc:      {recalc_t*1000:.2f}ms {'(daemon)' if is_daemon_running() else '(cold)'}"
        )
        print(f"   Double load: {reload_t*1000:.2f}ms")
        print(f"   TOTAL:       {total*1000:.2f}ms")
        print(f"\n   B1 computed value: {wb_values.active['B1'].value}")

        wb.close()
        wb_formula.close()
        wb_values.close()

        print("\n=== Summary ===")
        print(f"  openpyxl double-load overhead: {double_load_time*1000:.2f}ms")
        print(f"  Full sync cycle: {total*1000:.2f}ms")
        print("\n  Note: UNO approach would eliminate:")
        print("    - Double load (read formula + value from same object)")
        print("    - File I/O on recalc (in-memory calculateAll)")
        print("    - Save before recalc (modify in-memory)")

    finally:
        path.unlink(missing_ok=True)


if __name__ == "__main__":
    benchmark_openpyxl()
