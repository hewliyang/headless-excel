#!/usr/bin/env python3
"""Test openpyxl scaling with larger files."""

import tempfile
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook


def test_scale(num_rows: int, num_cols: int):
    """Test load times at different scales."""
    print(f"\n=== Testing {num_rows} rows x {num_cols} cols ===")

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = Path(f.name)

    try:
        # Create workbook with data
        print("Creating workbook...")
        t0 = time.perf_counter()
        wb = Workbook()
        ws = wb.active
        for row in range(1, num_rows + 1):
            for col in range(1, num_cols + 1):
                ws.cell(row=row, column=col, value=row * col)
        create_time = time.perf_counter() - t0
        print(f"  Create: {create_time:.2f}s")

        # Save
        t0 = time.perf_counter()
        wb.save(path)
        save_time = time.perf_counter() - t0
        wb.close()

        file_size = path.stat().st_size / (1024 * 1024)
        print(f"  Save: {save_time:.2f}s")
        print(f"  File size: {file_size:.2f} MB")

        # Load (formulas)
        t0 = time.perf_counter()
        wb = load_workbook(path)
        load_formula_time = time.perf_counter() - t0
        wb.close()
        print(f"  Load (formulas): {load_formula_time:.2f}s")

        # Load (data_only)
        t0 = time.perf_counter()
        wb = load_workbook(path, data_only=True)
        load_values_time = time.perf_counter() - t0
        wb.close()
        print(f"  Load (data_only): {load_values_time:.2f}s")

        # Double load (current sync)
        t0 = time.perf_counter()
        wb1 = load_workbook(path)
        wb2 = load_workbook(path, data_only=True)
        double_load_time = time.perf_counter() - t0
        wb1.close()
        wb2.close()
        print(f"  Double load: {double_load_time:.2f}s")

        return {
            "rows": num_rows,
            "cols": num_cols,
            "size_mb": file_size,
            "double_load_s": double_load_time,
        }

    finally:
        path.unlink(missing_ok=True)


if __name__ == "__main__":
    results = []

    # Test different scales
    for rows, cols in [(100, 10), (1000, 10), (5000, 20), (10000, 20)]:
        results.append(test_scale(rows, cols))

    print("\n=== Summary ===")
    print(f"{'Rows':>8} {'Cols':>6} {'Size MB':>10} {'Double Load':>12}")
    for r in results:
        print(
            f"{r['rows']:>8} {r['cols']:>6} {r['size_mb']:>10.2f} {r['double_load_s']:>12.2f}s"
        )
